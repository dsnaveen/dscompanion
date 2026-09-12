"""Step 1 (Load Data) service functions — the REST equivalent of
``dscompanion/app/interactive_step1.py``'s load/preview/duplicate-resolution logic, with every
``streamlit`` call stripped out. Ported deliberately, not reimplemented, so behaviour
(raw-header duplicate detection, position-based resolution) matches the UI exactly.
"""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path
from typing import Any

import pandas as pd

from dscompanion.api.config import api_settings
from dscompanion.api.schemas import (
    LoadDataConfirmRequest,
    LoadDataConfirmResponse,
    LoadDataPreviewRequest,
    LoadDataPreviewResponse,
)
from dscompanion.api.services.data_files import _DATA_ROOT
from dscompanion.api.state import RunState
from dscompanion.api.steps import next_step_id

logger = logging.getLogger(__name__)

__all__ = ["preview_load_data", "confirm_load_data"]


def _resolve_path(path: str) -> str:
    """Resolves a path relative to the shared data root, unless already absolute.

    ``GET /api/data-files`` returns paths relative to ``_DATA_ROOT`` (for display —
    mirrors the Streamlit picker's short labels); this makes those same strings work
    unchanged when passed straight back into preview/confirm. Absolute paths (e.g. an
    upload staged elsewhere, once Phase H.2 lands) pass through untouched.

    Args:
        path (str): Path as received from the client.

    Returns:
        str: Absolute path to read.
    """
    candidate = Path(path)
    if candidate.is_absolute():
        return str(candidate)
    return str(_DATA_ROOT / candidate)


def _read_raw_header(path: str, fmt: str, sheet_name: str | None) -> list[str]:
    """Read a file's literal header row, before pandas de-duplicates repeated names.

    Args:
        path (str): Path to the data file.
        fmt (str): ``"csv"``, ``"parquet"``, or ``"xlsx"``.
        sheet_name (str | None): Sheet name, required when ``fmt == "xlsx"``.

    Returns:
        list[str]: Raw header values in column order.
    """
    if fmt == "csv":
        with open(path, newline="", encoding="utf-8", errors="replace") as fh:
            first_line = fh.readline()
        try:
            delimiter = csv.Sniffer().sniff(first_line).delimiter
        except csv.Error:
            delimiter = ","
        return next(csv.reader([first_line], delimiter=delimiter))
    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        try:
            ws = wb[sheet_name] if sheet_name else wb.active
            row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        finally:
            wb.close()
        return [str(v) for v in row]
    return list(pd.read_parquet(path).columns)


def _load_dataframe(path: str, fmt: str, sheet_name: str | None) -> pd.DataFrame:
    """Read the full dataframe for the chosen path/format/sheet.

    Args:
        path (str): Path to the data file.
        fmt (str): ``"csv"``, ``"parquet"``, or ``"xlsx"``.
        sheet_name (str | None): Sheet name, used only when ``fmt == "xlsx"``.

    Returns:
        pd.DataFrame: The loaded data, as pandas parses it.
    """
    if fmt == "csv":
        return pd.read_csv(path)
    if fmt == "xlsx":
        return pd.read_excel(path, sheet_name=sheet_name)
    return pd.read_parquet(path)


def _find_duplicate_names(raw_header: list[str]) -> list[str]:
    """Find column names that appear more than once in a raw header.

    Args:
        raw_header (list[str]): Raw header values in column order.

    Returns:
        list[str]: Distinct names that occur 2+ times, in first-seen order.
    """
    counts: dict[str, int] = {}
    for name in raw_header:
        counts[name] = counts.get(name, 0) + 1
    return [name for name, count in counts.items() if count > 1]


def _apply_duplicate_resolution(
    df: pd.DataFrame, raw_header: list[str], resolutions: dict[str, str]
) -> pd.DataFrame:
    """Apply confirmed rename/drop choices for duplicate columns.

    ``resolutions`` keys are ``"{name}__{occurrence_index}"`` (0-based, matching the
    order occurrences appear in ``raw_header``); values are ``"rename:{new_name}"`` or
    ``"drop"`` — the wire encoding for ``LoadDataConfirmRequest.duplicate_resolutions``.

    Args:
        df (pd.DataFrame): Loaded dataframe (pandas-mangled column names).
        raw_header (list[str]): Raw header values in column order.
        resolutions (dict[str, str]): Per-occurrence resolution, keyed as above.

    Returns:
        pd.DataFrame: Copy of ``df`` with duplicate columns renamed/dropped.

    Raises:
        ValueError: If a resolution value isn't ``"drop"`` or ``"rename:<name>"``.
    """
    occurrence_positions: dict[str, list[int]] = {}
    for idx, name in enumerate(raw_header):
        occurrence_positions.setdefault(name, []).append(idx)

    new_columns = list(raw_header)
    drop_positions: set[int] = set()
    for key, action in resolutions.items():
        name, _, occ_str = key.rpartition("__")
        positions = occurrence_positions.get(name, [])
        occ_idx = int(occ_str)
        if occ_idx >= len(positions):
            continue
        col_idx = positions[occ_idx]
        if action == "drop":
            drop_positions.add(col_idx)
        elif action.startswith("rename:"):
            new_columns[col_idx] = action.removeprefix("rename:")
        else:
            raise ValueError(f"Invalid duplicate resolution action: {action!r}")

    keep_positions = [i for i in range(len(new_columns)) if i not in drop_positions]
    resolved = df.iloc[:, keep_positions].copy()
    resolved.columns = [new_columns[i] for i in keep_positions]
    return resolved


def _sample_rows(df: pd.DataFrame, n: int) -> list[dict[str, Any]]:
    """Converts the first ``n`` rows to JSON-safe records.

    Uses pandas' own JSON encoder (via ``to_json``) rather than ``to_dict`` —
    ``to_dict`` leaves ``NaN``/``NaT``/``Timestamp`` values as non-JSON-safe Python
    objects that FastAPI's response serialisation would choke on; ``to_json`` handles
    all three correctly (``NaN``/``NaT`` become ``null``, ``Timestamp`` becomes an ISO
    string).

    Args:
        df (pd.DataFrame): Loaded dataframe.
        n (int): Number of leading rows to include.

    Returns:
        list[dict[str, Any]]: JSON-safe row records, in column order.
    """
    return json.loads(df.head(n).to_json(orient="records", date_format="iso"))


def preview_load_data(request: LoadDataPreviewRequest) -> LoadDataPreviewResponse:
    """Loads the requested file and reports its shape/dtypes/duplicate columns, without
    persisting anything to the run.

    Args:
        request (LoadDataPreviewRequest): Path, format, and sheet name to load.

    Returns:
        LoadDataPreviewResponse: Columns, dtypes, row count, and any duplicate column
        names found in the raw header.
    """
    path = _resolve_path(request.path)
    raw_header = _read_raw_header(path, request.format, request.sheet_name)
    df = _load_dataframe(path, request.format, request.sheet_name)
    logger.info("Previewed load: %d rows x %d columns", len(df), len(df.columns))
    return LoadDataPreviewResponse(
        columns=list(df.columns),
        dtypes={col: str(dtype) for col, dtype in df.dtypes.items()},
        n_rows=len(df),
        duplicate_columns=_find_duplicate_names(raw_header),
        sample_rows=_sample_rows(df, api_settings.preview_sample_rows),
    )


def confirm_load_data(run: RunState, request: LoadDataConfirmRequest) -> LoadDataConfirmResponse:
    """Loads the requested file, applies duplicate-column resolutions, and persists the
    result as this run's working dataframe.

    Args:
        run (RunState): The run to persist into.
        request (LoadDataConfirmRequest): Path, format, sheet name, and duplicate
            resolutions.

    Returns:
        LoadDataConfirmResponse: The final column list and row count after resolution.
    """
    path = _resolve_path(request.path)
    raw_header = _read_raw_header(path, request.format, request.sheet_name)
    df = _load_dataframe(path, request.format, request.sheet_name)
    dup_names = _find_duplicate_names(raw_header)
    if dup_names:
        df = _apply_duplicate_resolution(df, raw_header, request.duplicate_resolutions)

    run.artifacts["df"] = df
    run.step_data["load_data"] = {
        "path": request.path,
        "format": request.format,
        "columns": list(df.columns),
        "n_rows": len(df),
    }
    run.step_confirmed["load_data"] = True
    run.step_status["load_data"] = "done"
    nxt = next_step_id("load_data")
    if nxt is not None:
        run.step_status[nxt] = "current"
    run.audit_trail.append(
        ("load_data", f"Loaded {Path(path).name} — {len(df):,} rows x {len(df.columns)} columns.")
    )
    logger.info(
        "run_id=%s confirmed Step 1 load: %d rows x %d columns",
        run.run_id,
        len(df),
        len(df.columns),
    )
    return LoadDataConfirmResponse(confirmed=True, columns=list(df.columns), n_rows=len(df))
