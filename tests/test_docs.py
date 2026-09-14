"""Tests for dscompanion.docs — ModelCard."""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from dscompanion.docs import ModelCard
from dscompanion.models import ModelFactory


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401

        return True
    except ImportError:
        return False


def _xlsxwriter_available() -> bool:
    try:
        import xlsxwriter  # noqa: F401

        return True
    except ImportError:
        return False


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


@pytest.fixture
def fitted_model(data_split):
    X = data_split.X_train.select_dtypes(include="number").fillna(0)
    y = data_split.y_train
    model = ModelFactory.build("classification", "logistic")
    model.fit(X, y)
    # Patch split.X_train with numeric version so evaluate() works
    from unittest.mock import MagicMock

    split = MagicMock()
    split.X_train = X
    split.y_train = y
    split.X_val = None
    split.X_test = data_split.X_test.select_dtypes(include="number").fillna(0)
    split.y_test = data_split.y_test
    split.X_oot = data_split.X_oot.select_dtypes(include="number").fillna(0)
    split.y_oot = data_split.y_oot
    return model, split


@pytest.fixture
def fitted_model_with_val(data_split):
    """Same as fitted_model, but with a real validation split populated —
    needed to test the Performance Metrics sheet's Validation section."""
    X = data_split.X_train.select_dtypes(include="number").fillna(0)
    y = data_split.y_train
    model = ModelFactory.build("classification", "logistic")
    model.fit(X, y)
    from unittest.mock import MagicMock

    split = MagicMock()
    split.X_train = X
    split.y_train = y
    split.X_val = data_split.X_val.select_dtypes(include="number").fillna(0)
    split.y_val = data_split.y_val
    split.X_test = data_split.X_test.select_dtypes(include="number").fillna(0)
    split.y_test = data_split.y_test
    split.X_oot = data_split.X_oot.select_dtypes(include="number").fillna(0)
    split.y_oot = data_split.y_oot
    return model, split


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


class TestModelCard:
    def test_generate_completes(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(
            model=model,
            split=split,
            author="Test Author",
            use_case="Unit test",
        )
        card.generate()
        assert len(card.sections_) == 14

    def test_model_summary_includes_package_versions(self, fitted_model):
        """R.3: 'Model Summary' gains a package_versions sub-dict for future
        debugging when versions differ across machines."""
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        package_versions = card.sections_["model_summary"]["package_versions"]
        assert package_versions["pandas"] == pd.__version__
        assert package_versions["numpy"] == np.__version__
        # xgboost isn't installed in this fixture's dependency set by
        # assumption — either a real version string or the explicit
        # "not installed" marker, never silently missing from the dict.
        assert "xgboost" in package_versions

    def test_key_package_versions_marks_missing_packages_explicitly(self):
        from dscompanion.docs.model_card import ModelCard

        versions = ModelCard._key_package_versions()
        for name, value in versions.items():
            assert value == "not installed" or value, f"{name} had an empty version string"

    def test_to_dict_has_all_section_keys(self, fitted_model):
        from dscompanion.docs.model_card import _SECTION_KEYS

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        d = card.to_dict()
        for key in _SECTION_KEYS:
            assert key in d, f"Missing section key: {key}"

    def test_to_json_writes_file_matching_to_dict(self, fitted_model, tmp_path):
        import json

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_json(tmp_path / "model_card.json")
        assert out.exists()
        with open(out, encoding="utf-8") as f:
            file_content = f.read()
        # Compare serialized strings, not loaded objects — some sections (e.g.
        # a PSI value on the first split) legitimately contain NaN, and
        # float('nan') != float('nan') would make an object-level == fail
        # even when the content genuinely matches.
        assert file_content == json.dumps(card.to_dict(), indent=2)

    def test_to_json_creates_parent_directories(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_json(tmp_path / "nested" / "dir" / "model_card.json")
        assert out.exists()

    @pytest.mark.skipif(not _docx_available(), reason="python-docx not installed")
    def test_to_word_creates_file_above_1kb(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_word(tmp_path / "model_card.docx")
        assert out.exists()
        assert out.stat().st_size > 1024

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_creates_non_empty_file(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        assert out.exists()
        assert out.stat().st_size > 1024

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_creates_one_sheet_per_top_level_section_minus_eda(
        self, fitted_model, tmp_path
    ):
        import openpyxl

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        sheetnames = openpyxl.load_workbook(out).sheetnames
        for expected in [
            "Config Summary",
            "Model Summary",
            "Data Lineage",
            "Feature Inventory",
            "Performance Metrics",
            "Decile Table",
            "Stability",
            "Calibration",
            "Explainability",
            "Hyperparameter Tuning",
            "Leaderboard",
            "Limitations",
            "Governance",
            "Full Config",
        ]:
            assert expected in sheetnames

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_creates_eda_subsheets_when_eda_report_provided(
        self, fitted_model, data_split, tmp_path
    ):
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        sheetnames = openpyxl.load_workbook(out).sheetnames
        for expected in [
            "EDA - Overview",
            "EDA - Numeric",
            "EDA - Categorical",
            "EDA - IV Ranking",
            "EDA - Correlations",
            "EDA - Missing",
            "EDA - Interactions",
        ]:
            assert expected in sheetnames

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_numeric_sheet_has_one_section_per_split(
        self, fitted_model, data_split, tmp_path
    ):
        """R.4: 'EDA - Numeric' extended to per-split sections (Train/Test/OOT
        for this fixture — no validation split), mirroring Decile Table /
        Performance Metrics.
        """
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Numeric"]
        # Column B is the label/first-data column (R.1: content starts at B2).
        col_b_values = [
            c.value for row in ws.iter_rows() for c in row if c.value is not None and c.column == 2
        ]
        for label in ("Train", "Test", "OOT"):
            assert label in col_b_values
        assert "Validation" not in col_b_values
        assert col_b_values.count("feature") == 3  # each section's own table header

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_categorical_sheet_has_train_section(self, fitted_model, data_split, tmp_path):
        """R.5: 'EDA - Categorical' extended to per-split sections, same pattern
        as EDA - Numeric. `fitted_model`'s Test/OOT are numeric-only (stripped
        by the fixture for the plain logistic model it fits), so only Train
        has categorical data here — the per-split *mechanism* itself is
        covered directly by test_eda_summary_by_split_includes_every_split
        below."""
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Categorical"]
        col_b_values = [
            c.value for row in ws.iter_rows() for c in row if c.value is not None and c.column == 2
        ]
        assert "Train" in col_b_values
        assert col_b_values.count("feature") == 1

    def test_eda_summary_by_split_includes_every_split(self, fitted_model_with_val):
        """Unit-level check of the per-split mechanism (R.4/R.5/R.6) with
        categorical data present on every split — end-to-end fixtures above
        only exercise numeric splits, since their model is fit on numeric-only
        features."""
        model, split = fitted_model_with_val
        train_cat = pd.DataFrame({"feature": ["c1"], "cardinality": [3]})
        split.X_test = pd.DataFrame({"c1": ["a", "b", "a"]})
        split.X_val = pd.DataFrame({"c1": ["a", "b"]})
        split.X_oot = pd.DataFrame({"c1": ["a", "a", "b"]})
        card = ModelCard(model=model, split=split)
        by_split = card._eda_summary_by_split(train_cat, "categorical")
        assert set(by_split.keys()) == {"Train", "Test", "Validation", "OOT"}
        assert by_split["Train"] is train_cat
        for label in ("Test", "Validation", "OOT"):
            assert "c1" in by_split[label]["feature"].tolist()

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_missing_sheet_omits_row_level_sample_matrix(
        self, fitted_model, data_split, tmp_path
    ):
        """The row-level True/False sample matrix served no purpose and was
        removed. The sheet keeps only the missing-% by feature table + chart,
        one such section per split instead of train-only.
        """
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Missing"]

        # The old row-level matrix ran up to missing_matrix_max_rows (500) data
        # rows on its own; the new per-split sectioned layout (Train/Test/OOT for
        # this fixture) is comfortably under that even with 3 stacked sections.
        assert ws.max_row < 100
        cell_values = [c.value for row in ws.iter_rows() for c in row]
        # bool check by type, not `in` (Python's `0 == False`, so `False in [0, 0]`
        # is True even though those are plain ints from the missing_pct column)
        assert not any(isinstance(v, bool) for v in cell_values)

        # The missing-% table + chart must still be present, once per split.
        assert cell_values.count("missing_pct") == 3
        for label in ("Train", "Test", "OOT"):
            assert label in cell_values

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_omits_eda_subsheets_when_no_eda_report(self, fitted_model, tmp_path):
        import openpyxl

        model, split = fitted_model
        card = ModelCard(model=model, split=split, eda_report=None)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        sheetnames = openpyxl.load_workbook(out).sheetnames
        assert not any(name.startswith("EDA - ") for name in sheetnames)

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_respects_sample_rows_compliance_gate(
        self, fitted_model, data_split, tmp_path
    ):
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()  # gate left at default False
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        sheetnames = openpyxl.load_workbook(out).sheetnames
        assert "EDA - Sample Rows" not in sheetnames

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_correlation_sheet_has_conditional_format(
        self, fitted_model, data_split, tmp_path
    ):
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Correlations"]
        assert len(ws.conditional_formatting._cf_rules) > 0

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_decile_sheet_has_chart(self, fitted_model, tmp_path):
        import openpyxl

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Decile Table"]
        assert len(ws._charts) > 0

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_numeric_charts_sheet_created_when_eda_report_provided(
        self, fitted_model, data_split, tmp_path
    ):
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Numeric Charts"]
        assert len(ws._charts) > 0

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_categorical_charts_sheet_created_when_eda_report_provided(
        self, fitted_model, data_split, tmp_path
    ):
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Categorical Charts"]
        assert len(ws._charts) > 0

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_categorical_charts_sheet_has_percentage_column(
        self, fitted_model, data_split, tmp_path
    ):
        """R.7: value/count table on 'EDA - Categorical Charts' gains a pct column."""
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Categorical Charts"]
        cell_values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        assert "pct" in cell_values
        # Train appears (this fixture's Test/OOT are numeric-only, see
        # test_eda_summary_by_split_includes_every_split for the per-split
        # mechanism covered on synthetic categorical data).
        assert "Train" in cell_values

    def test_categorical_charts_pct_column_present_on_every_split(self, fitted_model_with_val):
        """R.7 unit check: the pct column and per-split sections appear even
        though fitted_model_with_val's own splits are numeric-only — force
        categorical data onto every split directly, and stub eda_report with
        a plain UnivariateAnalyser (the same object
        EDAReport.categorical_value_counts() delegates to), rather than
        threading a real target column through EDAReport.run_all()."""
        import io

        import openpyxl

        from dscompanion.eda import UnivariateAnalyser

        model, split = fitted_model_with_val
        split.X_test = pd.DataFrame({"c1": ["a", "b", "a", "a"]})
        split.X_val = pd.DataFrame({"c1": ["a", "b"]})
        split.X_oot = pd.DataFrame({"c1": ["a", "a", "b"]})
        from types import SimpleNamespace

        train_analyser = UnivariateAnalyser().fit(
            SimpleNamespace(train_X=pd.DataFrame({"c1": ["a", "b", "a"]}))
        )
        card = ModelCard(model=model, split=split, eda_report=train_analyser)
        card._excel_sheet_registry = []
        categorical_df = pd.DataFrame({"feature": ["c1"]})

        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="xlsxwriter") as writer:
            card._write_categorical_charts_excel_sheet(writer, categorical_df)
        wb = openpyxl.load_workbook(buf)
        ws = wb["EDA - Categorical Charts"]
        cell_values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        assert "pct" in cell_values
        for label in ("Train", "Test", "Validation", "OOT"):
            assert label in cell_values

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_interactions_sheet_has_images_not_raw_data(
        self, fitted_model, data_split, tmp_path
    ):
        """R.8: 'EDA - Interactions' no longer writes a raw per-row data table
        (previously scaled with dataset row count) — it embeds a pre-rendered
        image per numeric feature pair instead."""
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Interactions"]
        assert len(ws._images) > 0
        # data_split has 3,503 training rows; a raw per-row table would need
        # that many rows. The sheet's actual row usage should stay tiny —
        # bounded by the number of feature pairs, not dataset size.
        assert ws.max_row < 100

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_charts_false_omits_chart_only_sheets(
        self, fitted_model, data_split, tmp_path
    ):
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report, excel_charts=False)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        sheetnames = openpyxl.load_workbook(out).sheetnames
        for omitted in [
            "EDA - Interactions",
            "EDA - Numeric Charts",
            "EDA - Categorical Charts",
        ]:
            assert omitted not in sheetnames

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_charts_false_keeps_data_tables_without_charts(
        self, fitted_model, data_split, tmp_path
    ):
        import openpyxl

        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report, excel_charts=False)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        wb = openpyxl.load_workbook(out)
        assert "EDA - IV Ranking" in wb.sheetnames
        assert len(wb["EDA - IV Ranking"]._charts) == 0
        assert "Decile Table" in wb.sheetnames
        assert len(wb["Decile Table"]._charts) == 0

    def test_numeric_histogram_data_clips_extreme_values(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        series = pd.Series(list(range(100)) + [10_000])
        hist_df, _ = card._numeric_histogram_data(series)
        assert hist_df["bin_center"].max() < 10_000

    def test_numeric_histogram_data_includes_kde_for_varied_series(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        rng = np.random.RandomState(0)
        series = pd.Series(rng.randn(200))
        hist_df, has_kde = card._numeric_histogram_data(series)
        assert has_kde is True
        assert (hist_df["kde"] != 0).any()

    def test_numeric_histogram_data_no_kde_for_constant_series(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        series = pd.Series([5.0] * 10)
        hist_df, has_kde = card._numeric_histogram_data(series)
        assert has_kde is False
        assert (hist_df["kde"] == 0).all()

    def test_numeric_histogram_data_empty_for_short_series(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        hist_df, has_kde = card._numeric_histogram_data(pd.Series([1.0]))
        assert hist_df.empty
        assert has_kde is False

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_index_sheet_is_first_and_lists_every_other_sheet(
        self, fitted_model, tmp_path
    ):
        import openpyxl

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames[0] == "Index"
        ws = wb["Index"]
        # Every sheet's content starts at B2 — the Index sheet's own header
        # follows the same convention: row 2, column B.
        assert [ws["B2"].value, ws["C2"].value, ws["D2"].value] == [
            "Sheet",
            "Purpose",
            "Feedback",
        ]
        listed = {ws.cell(row=r, column=2).value for r in range(3, ws.max_row + 1)}
        assert listed == set(wb.sheetnames) - {"Index"}

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_index_hyperlinks_resolve_to_listed_sheet(self, fitted_model, tmp_path):
        import openpyxl

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Index"]
        sheet_name = ws["B3"].value
        assert ws["B3"].hyperlink.location == f"'{sheet_name}'!A1"

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_sheet_a1_links_back_to_index(self, fitted_model, tmp_path):
        import openpyxl

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Config Summary"]
        # The nav link sits on row 1 (reserved nav row), but still in column B
        # (R.1's blank-column-A margin applies here too).
        assert ws["B1"].value == "Back to Index"
        assert ws["B1"].hyperlink.location == "'Index'!A1"

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_gridlines_hidden_on_every_sheet(self, fitted_model, tmp_path):
        import openpyxl

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        wb = openpyxl.load_workbook(out)
        for name in wb.sheetnames:
            assert wb[name].sheet_view.showGridLines is False, name

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_with_tuner_writes_best_params_as_parameter_value_table(
        self, fitted_model, tmp_path
    ):
        # Regression test: best_params_ is a flat dict of hyperparameter scalars
        # (e.g. {"max_depth": 5, "learning_rate": 0.1}) — pd.DataFrame() on a
        # bare dict of scalars raises "If using all scalar values, you must
        # pass an index". A tuned pipeline run crashed at to_excel() the
        # first time hyperparameter_tuning's section actually had content
        # (untuned runs only ever hit the scalar-only
        # {"note": "No tuner provided."} fallback, never this path).
        from unittest.mock import MagicMock

        import openpyxl

        model, split = fitted_model
        tuner = MagicMock()
        tuner.backend = "optuna"
        tuner.n_trials = 10
        tuner.metric = "roc_auc"
        tuner.best_params_ = {"max_depth": 5, "learning_rate": 0.1, "n_estimators": 100}
        tuner.best_score_ = 0.87

        card = ModelCard(model=model, split=split, tuner=tuner)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")  # must not raise

        wb = openpyxl.load_workbook(out)
        ws = wb["Hyperparameter Tuning"]
        # Column A is a blank margin (R.1: content starts at B2) — strip the
        # leading None before comparing row shape.
        rows = [tuple(v for v in r if v is not None) for r in ws.iter_rows(values_only=True)]
        assert ("max_depth", 5) in rows
        assert ("learning_rate", 0.1) in rows
        assert ("n_estimators", 100) in rows

    def test_to_excel_raises_importerror_when_xlsxwriter_missing(
        self, fitted_model, tmp_path, monkeypatch
    ):
        import builtins

        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()

        real_import = builtins.__import__

        def _fake_import(name, *args, **kwargs):
            if name == "xlsxwriter":
                raise ImportError("simulated missing xlsxwriter")
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", _fake_import)
        with pytest.raises(ImportError, match="xlsxwriter"):
            card.to_excel(tmp_path / "model_card.xlsx")

    def test_no_config_checks_renders_empty_schema(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split, config_checks=None)
        card.generate()
        assert card.sections_["config_summary"].empty
        assert list(card.sections_["config_summary"].columns) == [
            "parameter",
            "choices",
            "default",
            "user_choice",
            "is_deviation",
        ]

    def test_config_checks_populate_section_dataframe(self, fitted_model):
        model, split = fitted_model
        checks = [
            {
                "parameter": "eda.enabled",
                "choices": "True / False",
                "default": True,
                "user_choice": False,
                "is_deviation": True,
            },
            {
                "parameter": "eda.univariate",
                "choices": "True / False",
                "default": True,
                "user_choice": True,
                "is_deviation": False,
            },
        ]
        card = ModelCard(model=model, split=split, config_checks=checks)
        card.generate()
        df = card.sections_["config_summary"]
        assert len(df) == 2
        assert df.iloc[0]["parameter"] == "eda.enabled"
        assert bool(df.iloc[0]["is_deviation"]) is True
        assert bool(df.iloc[1]["is_deviation"]) is False

    def test_to_html_creates_non_empty_file(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_html(tmp_path / "model_card.html")
        assert out.exists()
        assert out.stat().st_size > 500

    def test_to_html_includes_every_navbar_section_always_present(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        content = card.to_html(tmp_path / "card.html").read_text()
        for anchor in (
            "#overview",
            "#performance",
            "#feature-inventory",
            "#full-config",
            "#governance",
        ):
            assert f'href="{anchor}"' in content

    def test_decile_table_disabled_returns_empty(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split, decile_table=False)
        card.generate()
        df = card.sections_["decile_table"]
        assert df.empty
        assert list(df.columns) == [
            "decile",
            "count",
            "events",
            "event_rate",
            "cumulative_count",
            "cumulative_events",
            "cumulative_event_rate",
            "pct_of_total_events",
            "cumulative_pct_of_total_events",
            "lift",
            "cumulative_lift",
        ]

    def test_decile_table_not_applicable_for_regression(self, fitted_model):
        from dscompanion.models import ModelFactory

        model, split = fitted_model
        X = split.X_train.select_dtypes(include="number").fillna(0)
        reg_model = ModelFactory.build("regression", "linear")
        reg_model.fit(X, X.iloc[:, 0])  # arbitrary numeric target

        card = ModelCard(model=reg_model, split=split)
        card.generate()
        assert card.sections_["decile_table"].empty

    def test_to_html_omits_eda_navbar_sections_when_no_eda_report(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split, eda_report=None)
        card.generate()
        content = card.to_html(tmp_path / "card.html").read_text()
        assert 'href="#eda-variables"' not in content

    def test_to_html_includes_eda_navbar_sections_when_eda_report_provided(
        self, fitted_model, data_split, tmp_path
    ):
        from dscompanion.eda import EDAReport

        model, split = fitted_model
        eda_report = EDAReport(data_split, target="target").run_all()
        card = ModelCard(model=model, split=split, eda_report=eda_report)
        card.generate()
        content = card.to_html(tmp_path / "card.html").read_text()
        assert 'href="#eda-variables"' in content

    def test_to_html_shows_model_algorithm_in_overview(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        content = card.to_html(tmp_path / "card.html").read_text()
        assert type(model.estimator).__name__ in content

    def test_to_html_creates_parent_directories(self, fitted_model, tmp_path):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        nested = tmp_path / "a" / "b" / "card.html"
        out = card.to_html(nested)
        assert out.exists()

    def test_to_html_package_versions_render_as_own_table_not_repr(self, fitted_model, tmp_path):
        """R.3: package_versions must not be dumped via the generic
        stat_row(v) loop — that would render Python's repr() of the dict
        into a single table cell instead of proper rows."""
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        content = card.to_html(tmp_path / "card.html").read_text()
        assert "Package versions" in content
        assert "pandas" in content
        assert pd.__version__ in content
        # A dict repr in the page would look like "{'pandas':" — must not appear.
        assert "{'pandas'" not in content

    def test_to_tracking_artifact_logs_excel_html_and_json(self, fitted_model, monkeypatch):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()

        logged_paths = []
        monkeypatch.setattr(
            "dscompanion.tracking.log_artifact",
            lambda path, artifact_path=None: logged_paths.append(path),
        )

        card.to_tracking_artifact()

        assert any(p.endswith(".xlsx") for p in logged_paths)
        assert any(p.endswith(".html") for p in logged_paths)
        assert any(p.endswith(".json") for p in logged_paths)


def test_numeric_histogram_data_delegates_to_shared_widget(fitted_model):
    """ModelCard._numeric_histogram_data must stay behavior-identical after
    delegating to html_widgets.compute_histogram (DRY refactor, Task 7)."""
    import pandas as pd

    from dscompanion.config import settings
    from dscompanion.docs import html_widgets as hw

    model, split = fitted_model
    card = ModelCard(model=model, split=split)
    series = pd.Series(range(100), dtype=float)
    hist_df, has_kde = card._numeric_histogram_data(series)
    expected_df, expected_kde = hw.compute_histogram(
        series,
        bins=settings.eda_histogram_bins,
        clip_lower_pct=card.excel_chart_clip_lower_pct,
        clip_upper_pct=card.excel_chart_clip_upper_pct,
    )
    pd.testing.assert_frame_equal(
        hist_df.reset_index(drop=True), expected_df.reset_index(drop=True)
    )
    assert has_kde == expected_kde


class _StubSHAPExplainer:
    def mean_abs_shap(self):
        return pd.DataFrame({"feature": ["a", "b"], "mean_abs_shap": [0.5, 0.2], "rank": [1, 2]})


class _StubPermutationImportance:
    def importance_table(self):
        return pd.DataFrame(
            {
                "feature": ["a", "b"],
                "importance_mean": [0.4, 0.1],
                "importance_std": [0.01, 0.02],
                "rank": [1, 2],
            }
        )


class TestModelCardExplainabilitySection:
    """PI.4 — _explainability() independently reports SHAP and permutation importance."""

    def test_note_when_neither_provided(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        result = card._explainability()
        assert result == {"note": "No SHAP explainer or permutation importance provided."}

    def test_shap_only_unchanged(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split, explainer=_StubSHAPExplainer())
        result = card._explainability()
        assert "top_features" in result
        assert "permutation_top_features" not in result
        assert result["top_features"][0]["feature"] == "a"

    def test_permutation_only(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(
            model=model, split=split, permutation_importance=_StubPermutationImportance()
        )
        result = card._explainability()
        assert "permutation_top_features" in result
        assert "top_features" not in result
        assert result["permutation_top_features"][0]["feature"] == "a"

    def test_both_provided(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(
            model=model,
            split=split,
            explainer=_StubSHAPExplainer(),
            permutation_importance=_StubPermutationImportance(),
        )
        result = card._explainability()
        assert "top_features" in result
        assert "permutation_top_features" in result


class TestPerformanceMetricsByLabelledSection:
    """Excel-only Performance Metrics sheet: wide Train/Test/Validation/OOT
    comparison table plus diagnostics — Word/HTML/to_dict() keep the
    original wide, train-test-OOT-only shape via _performance_metrics(),
    unchanged."""

    def test_original_wide_metrics_unchanged_by_the_new_method(self, fitted_model_with_val):
        """_performance_metrics() (Word/HTML/to_dict()) must still exclude
        validation and stay a single wide table — the new by-split method
        is additive, not a replacement."""
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        wide = card._performance_metrics()
        assert isinstance(wide, pd.DataFrame)
        assert "Validation" not in wide.columns
        assert "val" not in wide.columns

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_excel_sheet_comparison_table_has_all_splits(self, fitted_model_with_val, tmp_path):
        import openpyxl

        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Performance Metrics"]
        cell_values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        for label in ("Train", "Test", "Validation", "OOT"):
            assert label in cell_values  # comparison table column headers
        assert "metric" in cell_values  # comparison table's own header
        assert "roc_auc" in cell_values
        assert "Diagnostics" in cell_values
        # No more per-split stacked sections (R.9): the per-split labels
        # each appear exactly once — as the comparison table's header row —
        # not a second time as a stacked section's own label row.
        assert cell_values.count("Train") == 1

    def test_comparison_table_includes_all_four_splits(self, fitted_model_with_val):
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        wide = card._comparison_metrics_table()
        for col in ("Train", "Test", "Validation", "OOT"):
            assert col in wide.columns
        assert "metric" in wide.columns

    def test_comparison_table_omits_missing_split(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        wide = card._comparison_metrics_table()
        assert "Validation" not in wide.columns
        assert {"Train", "Test", "OOT"} <= set(wide.columns)

    def test_diagnostics_flags_overfitting_gap(self, fitted_model_with_val):
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        crafted = pd.DataFrame(
            [
                {"split": "train", "metric": "roc_auc", "value": 0.95},
                {"split": "test", "metric": "roc_auc", "value": 0.60},
            ]
        )
        card.model.evaluate = lambda _split: crafted
        flags = card._performance_diagnostics()
        assert any("overfit" in f["message"].lower() for f in flags)

    def test_diagnostics_flags_leakage(self, fitted_model_with_val):
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        crafted = pd.DataFrame(
            [
                {"split": "train", "metric": "roc_auc", "value": 0.995},
                {"split": "test", "metric": "roc_auc", "value": 0.995},
            ]
        )
        card.model.evaluate = lambda _split: crafted
        flags = card._performance_diagnostics()
        assert any("leak" in f["message"].lower() for f in flags)

    def test_diagnostics_empty_when_metrics_are_healthy(self, fitted_model_with_val):
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        crafted = pd.DataFrame(
            [
                {"split": "train", "metric": "roc_auc", "value": 0.75},
                {"split": "test", "metric": "roc_auc", "value": 0.72},
            ]
        )
        card.model.evaluate = lambda _split: crafted
        assert card._performance_diagnostics() == []

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_excel_sheet_has_comparison_table_above_diagnostics(
        self, fitted_model_with_val, tmp_path
    ):
        import openpyxl

        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Performance Metrics"]
        # Section/comparison/diagnostics labels are always written to column B
        # (R.1: content starts at B2, column A is a blank margin); the
        # comparison table's own "Train"/etc. column headers are NOT in
        # column B (they start there but run rightward).
        col_b_values = {
            c.value for row in ws.iter_rows() for c in row if c.value is not None and c.column == 2
        }
        assert "Comparison (all splits)" in col_b_values
        assert "Diagnostics" in col_b_values

        def _row_of(value: str) -> int:
            return next(
                c.row for row in ws.iter_rows() for c in row if c.value == value and c.column == 2
            )

        assert _row_of("Comparison (all splits)") < _row_of("Diagnostics")


class TestDecileTableByLabelledSection:
    """Excel-only sectioned Decile Table sheet (Train/Test/Validation/OOT),
    each with its own chart — mirrors the Performance Metrics sheet's
    sectioning. Word/HTML/to_dict() keep _decile_table_section()'s original
    test-only table, unchanged."""

    def test_by_split_includes_validation_when_present(self, fitted_model_with_val):
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        by_split = card._decile_table_by_split()
        assert set(by_split.keys()) == {"Train", "Test", "Validation", "OOT"}
        for df in by_split.values():
            assert len(df) > 0

    def test_by_split_omits_validation_when_absent(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        by_split = card._decile_table_by_split()
        assert "Validation" not in by_split
        assert set(by_split.keys()) == {"Train", "Test", "OOT"}

    def test_by_split_empty_when_decile_table_disabled(self, fitted_model_with_val):
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split, decile_table=False)
        assert card._decile_table_by_split() == {}

    def test_original_test_only_decile_unchanged_by_the_new_method(self, fitted_model_with_val):
        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        original = card._decile_table_section()
        assert isinstance(original, pd.DataFrame)
        assert len(original) > 0  # test split only, unchanged shape

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_excel_sheet_has_one_labelled_chart_per_split(self, fitted_model_with_val, tmp_path):
        import openpyxl

        model, split = fitted_model_with_val
        card = ModelCard(model=model, split=split)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Decile Table"]
        cell_values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        for label in ("Train", "Test", "Validation", "OOT"):
            assert label in cell_values
        assert len(ws._charts) == 4


class TestLeaderboardSection:
    def test_no_leaderboard_provided_returns_empty_dataframe(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        card.generate()
        section = card.sections_["leaderboard"]
        assert isinstance(section, pd.DataFrame)
        assert section.empty
        assert "algorithm" in section.columns

    def test_leaderboard_provided_populates_section_with_ranked_table(self, fitted_model):
        from dscompanion.leaderboard import Leaderboard

        model, split = fitted_model
        lb = Leaderboard(include=["logistic", "decision_tree"])
        lb.run(split)

        card = ModelCard(model=model, split=split, leaderboard=lb)
        card.generate()
        section = card.sections_["leaderboard"]
        assert set(section["algorithm"]) == {"logistic", "decision_tree"}

    def test_to_dict_includes_leaderboard_records(self, fitted_model):
        from dscompanion.leaderboard import Leaderboard

        model, split = fitted_model
        lb = Leaderboard(include=["logistic"])
        lb.run(split)

        card = ModelCard(model=model, split=split, leaderboard=lb)
        card.generate()
        d = card.to_dict()
        assert d["leaderboard"][0]["algorithm"] == "logistic"

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_leaderboard_sheet_has_ranked_rows(self, fitted_model, tmp_path):
        import openpyxl

        from dscompanion.leaderboard import Leaderboard

        model, split = fitted_model
        lb = Leaderboard(include=["logistic", "decision_tree"])
        lb.run(split)

        card = ModelCard(model=model, split=split, leaderboard=lb)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        sheet = openpyxl.load_workbook(out)["Leaderboard"]
        # Column A is a blank margin (R.1: content starts at B2); the
        # "algorithm" column is B. Bounded to the ranking table's own 2
        # data rows — R.10 appends a second, per-split "algorithm" table
        # further down the same sheet.
        algorithms = {
            row[0].value for row in sheet.iter_rows(min_row=3, max_row=4, min_col=2, max_col=2)
        }
        assert algorithms == {"logistic", "decision_tree"}

    @pytest.mark.skipif(not _xlsxwriter_available(), reason="xlsxwriter not installed")
    def test_to_excel_leaderboard_sheet_has_per_split_metrics_table(self, fitted_model, tmp_path):
        """R.10: 'do the same for all splits' for Leaderboard, scoped Excel-only —
        a second table below the (unchanged) ranking table, one row per
        algorithm per split."""
        import openpyxl

        from dscompanion.leaderboard import Leaderboard

        model, split = fitted_model
        lb = Leaderboard(include=["logistic", "decision_tree"])
        lb.run(split)

        card = ModelCard(model=model, split=split, leaderboard=lb)
        card.generate()
        out = card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Leaderboard"]
        cell_values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        assert "Per-Algorithm Metrics — All Splits" in cell_values
        # fitted_model has no validation split.
        for label in ("Train", "Test", "OOT"):
            assert label in cell_values
        assert "Validation" not in cell_values
        # The per-split table's own "algorithm" header, distinct from the
        # ranking table's "algorithm" column header above it.
        assert cell_values.count("algorithm") == 4  # 1 ranking table + 3 split sections

    def test_leaderboard_metrics_by_split_reuses_fitted_models(self, fitted_model_with_val):
        """Unit check: _leaderboard_metrics_by_split() re-evaluates
        Leaderboard.fitted_models_ across every split without refitting."""
        from dscompanion.leaderboard import Leaderboard

        model, split = fitted_model_with_val
        lb = Leaderboard(include=["logistic"])
        lb.run(split)

        card = ModelCard(model=model, split=split, leaderboard=lb)
        by_split = card._leaderboard_metrics_by_split()
        assert set(by_split.keys()) == {"Train", "Test", "Validation", "OOT"}
        for split_df in by_split.values():
            assert list(split_df["algorithm"]) == ["logistic"]
            assert "roc_auc" in split_df.columns

    def test_leaderboard_metrics_by_split_empty_when_no_leaderboard(self, fitted_model):
        model, split = fitted_model
        card = ModelCard(model=model, split=split)
        assert card._leaderboard_metrics_by_split() == {}
