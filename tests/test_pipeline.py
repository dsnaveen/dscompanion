"""Tests for dscompanion.pipeline — PipelineConfig validation and PipelineRunner orchestration."""

from __future__ import annotations

import logging
import re
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import yaml
from pydantic import ValidationError

from dscompanion.features import ColumnRecipe, FeatureTransformChain, TransformStep
from dscompanion.pipeline.config import (
    DataConfig,
    DistributionConfig,
    EDAConfig,
    EncoderConfig,
    ExplainConfig,
    FeaturesConfig,
    ImbalanceConfig,
    ImputerConfig,
    LeaderboardConfig,
    ModelConfig,
    PipelineConfig,
    ReportingConfig,
    ScalerConfig,
    SelectionConfig,
    SplitConfig,
    TargetConfig,
    TuningConfig,
)
from dscompanion.pipeline.runner import PipelineRunner


def _minimal_config(**overrides: object) -> PipelineConfig:
    """Build a minimal valid PipelineConfig, with top-level field overrides applied."""
    base = {
        "name": "test_experiment",
        "data": DataConfig(path="dummy.parquet", target="y"),
        "model": ModelConfig(task="classification"),
    }
    base.update(overrides)
    return PipelineConfig(**base)


# ---------------------------------------------------------------------------
# PipelineConfig validation
# ---------------------------------------------------------------------------


class TestPipelineConfigValidation:
    def test_minimal_config_constructs(self):
        cfg = _minimal_config()
        assert cfg.name == "test_experiment"
        assert cfg.data.target == "y"
        assert cfg.model.task == "classification"

    def test_reporting_config_has_no_mlflow_fields(self):
        cfg = _minimal_config()
        assert not hasattr(cfg.reporting, "mlflow_experiment")
        assert not hasattr(cfg.reporting, "mlflow_enabled")

    def test_missing_name_raises(self):
        with pytest.raises(ValidationError):
            PipelineConfig(
                data=DataConfig(path="dummy.parquet", target="y"),
                model=ModelConfig(task="classification"),
            )

    def test_missing_data_raises(self):
        with pytest.raises(ValidationError):
            PipelineConfig(name="exp", model=ModelConfig(task="classification"))

    def test_missing_model_raises(self):
        with pytest.raises(ValidationError):
            PipelineConfig(name="exp", data=DataConfig(path="dummy.parquet", target="y"))

    def test_unknown_top_level_field_rejected(self):
        with pytest.raises(ValidationError, match="extra_forbidden|Extra inputs"):
            _minimal_config(not_a_real_field=True)

    def test_invalid_data_format_raises(self):
        with pytest.raises(ValidationError, match="format must be one of"):
            DataConfig(path="dummy.parquet", target="y", format="xml")

    def test_feature_columns_and_ignore_columns_together_raises(self):
        with pytest.raises(ValidationError, match="mutually exclusive"):
            DataConfig(
                path="dummy.parquet",
                target="y",
                feature_columns=["a"],
                ignore_columns=["b"],
            )

    def test_ignore_columns_containing_target_raises(self):
        with pytest.raises(ValidationError, match="cannot include the target or date_column"):
            DataConfig(path="dummy.parquet", target="y", ignore_columns=["y"])

    def test_ignore_columns_containing_date_column_raises(self):
        with pytest.raises(ValidationError, match="cannot include the target or date_column"):
            DataConfig(
                path="dummy.parquet",
                target="y",
                date_column="ts",
                ignore_columns=["ts"],
            )

    def test_invalid_split_method_raises(self):
        with pytest.raises(ValidationError, match="method must be one of"):
            SplitConfig(method="bogus")

    def test_split_fraction_out_of_range_raises(self):
        with pytest.raises(ValidationError, match=r"Must be in \(0, 1\)"):
            SplitConfig(test_size=1.5)

    def test_invalid_model_task_raises(self):
        with pytest.raises(ValidationError, match="task must be one of"):
            ModelConfig(task="forecasting")

    def test_invalid_model_algorithm_raises(self):
        with pytest.raises(ValidationError, match="algorithm must be one of"):
            ModelConfig(task="classification", algorithm="prophet")

    @pytest.mark.parametrize(
        "task,algorithm",
        [
            ("classification", "svm"),
            ("classification", "knn"),
            ("classification", "decision_tree"),
            ("classification", "extra_trees"),
            ("classification", "adaboost"),
            ("classification", "naive_bayes"),
            ("regression", "elastic_net"),
            ("regression", "gradient_boosting"),
            ("regression", "svm"),
            ("regression", "knn"),
            ("regression", "decision_tree"),
            ("regression", "extra_trees"),
            ("regression", "adaboost"),
            ("clustering", "hierarchical"),
        ],
    )
    def test_model_algorithm_validates_against_model_factory(self, task, algorithm):
        # ModelConfig.valid_algorithm derives its allowed set from
        # ModelFactory.SUPPORTED_ALGORITHMS[task] rather than a separately
        # maintained literal — this covers every algorithm ModelFactory
        # supports, including "hierarchical" clustering, which a stale
        # hardcoded set previously rejected even though ModelFactory has
        # always supported it.
        cfg = ModelConfig(task=task, algorithm=algorithm)
        assert cfg.algorithm == algorithm

    def test_model_algorithm_error_names_task_specific_choices(self):
        with pytest.raises(ValidationError, match="for task='regression'"):
            ModelConfig(task="regression", algorithm="naive_bayes")

    def test_negative_iv_threshold_raises(self):
        with pytest.raises(ValidationError, match="iv_threshold must be >= 0"):
            SelectionConfig(iv_threshold=-0.1)

    def test_model_task_is_lowercased(self):
        cfg = ModelConfig(task="CLASSIFICATION")
        assert cfg.task == "classification"

    def test_eda_config_default_thresholds(self):
        cfg = EDAConfig()
        assert cfg.skewness_alert_threshold == 1.0
        assert cfg.zero_pct_alert_threshold == 0.05
        assert cfg.imbalance_alert_threshold == 0.5
        assert cfg.extreme_values_n == 5
        assert cfg.near_zero_variance_threshold == 0.01

    def test_eda_config_negative_skewness_threshold_raises(self):
        with pytest.raises(ValidationError, match="skewness_alert_threshold must be >= 0"):
            EDAConfig(skewness_alert_threshold=-1.0)

    def test_eda_config_negative_near_zero_variance_threshold_raises(self):
        with pytest.raises(ValidationError, match="near_zero_variance_threshold must be >= 0"):
            EDAConfig(near_zero_variance_threshold=-0.01)

    @pytest.mark.parametrize("field", ["zero_pct_alert_threshold", "imbalance_alert_threshold"])
    def test_eda_config_threshold_out_of_01_raises(self, field):
        with pytest.raises(ValidationError, match=r"Threshold must be in \(0, 1\)"):
            EDAConfig(**{field: 1.5})

    def test_eda_config_extreme_values_n_must_be_positive(self):
        with pytest.raises(ValidationError, match="extreme_values_n must be >= 1"):
            EDAConfig(extreme_values_n=0)

    def test_eda_config_thresholds_overridable_via_yaml(self):
        cfg = _minimal_config(eda=EDAConfig(skewness_alert_threshold=2.5, extreme_values_n=10))
        assert cfg.eda.skewness_alert_threshold == 2.5
        assert cfg.eda.extreme_values_n == 10

    def test_imputer_config_column_strategies_default_empty(self):
        cfg = ImputerConfig()
        assert cfg.column_strategies == {}
        assert cfg.column_fill_values == {}

    def test_imputer_config_invalid_column_strategy_raises(self):
        with pytest.raises(ValidationError, match="column_strategies has invalid values"):
            ImputerConfig(column_strategies={"col_a": "bogus"})

    @pytest.mark.parametrize("strategy", ["auto", "mean", "median", "constant", "most_frequent"])
    def test_imputer_config_valid_column_strategies_accepted(self, strategy):
        cfg = ImputerConfig(column_strategies={"col_a": strategy})
        assert cfg.column_strategies == {"col_a": strategy}


# ---------------------------------------------------------------------------
# ExplainConfig — permutation importance fields (PI.2)
# ---------------------------------------------------------------------------


class TestExplainConfigValidation:
    def test_permutation_fields_default(self):
        cfg = ExplainConfig()
        assert cfg.permutation_enabled is False
        assert cfg.permutation_n_repeats == 10
        assert cfg.permutation_sample_size == 5000
        assert cfg.permutation_top_n == 20

    def test_permutation_fields_override(self):
        cfg = ExplainConfig(
            permutation_enabled=True,
            permutation_n_repeats=5,
            permutation_sample_size=1000,
            permutation_top_n=10,
        )
        assert cfg.permutation_enabled is True
        assert cfg.permutation_n_repeats == 5
        assert cfg.permutation_sample_size == 1000
        assert cfg.permutation_top_n == 10

    def test_unknown_explain_field_rejected(self):
        with pytest.raises(ValidationError, match="extra_forbidden|Extra inputs"):
            ExplainConfig(permutation_enable_typo=True)


# ---------------------------------------------------------------------------
# Pre-flight error / auto-correction cases
# ---------------------------------------------------------------------------


class TestPreflight:
    def test_temporal_without_date_column_raises(self):
        cfg = _minimal_config(split=SplitConfig(method="temporal"))
        runner = PipelineRunner(cfg)
        with pytest.raises(ValueError, match="split.method='temporal' requires data.date_column"):
            runner._preflight()

    def test_temporal_with_date_column_passes(self):
        cfg = _minimal_config(
            data=DataConfig(path="dummy.parquet", target="y", date_column="snapshot_date"),
            split=SplitConfig(method="temporal"),
        )
        runner = PipelineRunner(cfg)
        runner._preflight()  # should not raise

    def test_grouped_without_group_column_raises(self):
        cfg = _minimal_config(split=SplitConfig(method="grouped"))
        runner = PipelineRunner(cfg)
        with pytest.raises(ValueError, match="split.method='grouped' requires split.group_column"):
            runner._preflight()

    def test_grouped_with_group_column_passes(self):
        cfg = _minimal_config(split=SplitConfig(method="grouped", group_column="customer_id"))
        runner = PipelineRunner(cfg)
        runner._preflight()  # should not raise

    def test_logistic_without_scaler_auto_corrects(self, caplog):
        cfg = _minimal_config(
            model=ModelConfig(task="classification", algorithm="logistic"),
            features=FeaturesConfig(scaler=ScalerConfig(strategy="none")),
        )
        runner = PipelineRunner(cfg)
        with caplog.at_level(logging.WARNING):
            runner._preflight()
        assert cfg.features.scaler.strategy == "standard"
        assert any("Auto-correcting" in r.message for r in caplog.records)

    def test_logistic_with_scaler_already_set_not_overridden(self):
        cfg = _minimal_config(
            model=ModelConfig(task="classification", algorithm="logistic"),
            features=FeaturesConfig(scaler=ScalerConfig(strategy="robust")),
        )
        runner = PipelineRunner(cfg)
        runner._preflight()
        assert cfg.features.scaler.strategy == "robust"

    def test_winsorizer_disabled_by_default(self):
        cfg = _minimal_config()
        assert cfg.features.winsorizer.enabled is False
        assert cfg.features.winsorizer.lower_tail == 0.01
        assert cfg.features.winsorizer.upper_tail == 0.01

    def test_winsorizer_tail_fraction_out_of_range_raises(self):
        from dscompanion.pipeline.config import WinsorizerConfig

        with pytest.raises(ValidationError):
            WinsorizerConfig(lower_tail=0.6)
        with pytest.raises(ValidationError):
            WinsorizerConfig(upper_tail=-0.1)

    def test_low_n_trials_warns_but_does_not_raise(self, caplog):
        cfg = _minimal_config(tuning=TuningConfig(enabled=True, n_trials=2))
        runner = PipelineRunner(cfg)
        with caplog.at_level(logging.WARNING):
            runner._preflight()
        assert any("very low" in r.message for r in caplog.records)

    def test_imbalance_strategy_on_regression_warns(self, caplog):
        cfg = _minimal_config(
            model=ModelConfig(task="regression", algorithm="xgboost"),
            target=TargetConfig(),
        )
        runner = PipelineRunner(cfg)
        with caplog.at_level(logging.WARNING):
            runner._preflight()
        assert any("only applicable to classification" in r.message for r in caplog.records)

    def test_vif_enabled_warns(self, caplog):
        cfg = _minimal_config(selection=SelectionConfig(vif_enabled=True))
        runner = PipelineRunner(cfg)
        with caplog.at_level(logging.WARNING):
            runner._preflight()
        assert any("vif_enabled" in r.message for r in caplog.records)

    def test_default_config_preflight_clean(self, caplog):
        cfg = _minimal_config()
        runner = PipelineRunner(cfg)
        with caplog.at_level(logging.WARNING):
            runner._preflight()
        assert caplog.records == []


# ---------------------------------------------------------------------------
# run_id generation — yyyymmdd_hhmmss, collision-safe
# ---------------------------------------------------------------------------


class TestGenerateRunIdAndDir:
    def test_format_is_timestamp(self, tmp_path):
        runner = PipelineRunner(_minimal_config())
        run_id, run_dir = runner._generate_run_id_and_dir(tmp_path)
        assert re.fullmatch(r"\d{8}_\d{6}", run_id)
        assert run_dir == tmp_path / run_id

    def test_returned_dir_does_not_already_exist(self, tmp_path):
        runner = PipelineRunner(_minimal_config())
        _, run_dir = runner._generate_run_id_and_dir(tmp_path)
        assert not run_dir.exists()

    def test_same_second_collision_gets_disambiguated(self, tmp_path):
        runner = PipelineRunner(_minimal_config())
        run_id, run_dir = runner._generate_run_id_and_dir(tmp_path)
        run_dir.mkdir(parents=True)  # simulate an existing run from the same second

        run_id_2, run_dir_2 = runner._generate_run_id_and_dir(tmp_path)
        assert run_id_2 != run_id
        assert run_id_2.startswith(run_id + "_")
        assert not run_dir_2.exists()


# ---------------------------------------------------------------------------
# Config deviation detection
# ---------------------------------------------------------------------------


class TestConfigDeviationDetection:
    def test_default_config_has_no_deviations(self):
        cfg = _minimal_config()
        runner = PipelineRunner(cfg)
        assert runner._detect_config_deviations() == []

    def test_eda_disabled_is_flagged(self):
        cfg = _minimal_config(eda=EDAConfig(enabled=False))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "eda.enabled" for d in deviations)

    def test_leakage_check_disabled_is_flagged(self):
        cfg = _minimal_config(selection=SelectionConfig(leakage_check=False))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "selection.leakage_check" for d in deviations)

    def test_shap_enabled_is_flagged(self):
        # shap_enabled defaults to False (opt-in, since SHAP is relatively
        # expensive) — enabling it is the deviation from default, not disabling it.
        cfg = _minimal_config(explain=ExplainConfig(shap_enabled=True))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "explain.shap_enabled" for d in deviations)

    def test_leaderboard_enabled_is_flagged(self):
        cfg = _minimal_config(leaderboard=LeaderboardConfig(enabled=True))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "leaderboard.enabled" for d in deviations)

    def test_permutation_enabled_is_flagged(self):
        cfg = _minimal_config(explain=ExplainConfig(permutation_enabled=True))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "explain.permutation_enabled" for d in deviations)


class TestLeaderboardConfigValidation:
    def test_disabled_by_default(self):
        cfg = _minimal_config()
        assert cfg.leaderboard.enabled is False

    def test_enabled_with_classification_task_constructs(self):
        cfg = _minimal_config(
            model=ModelConfig(task="classification"),
            leaderboard=LeaderboardConfig(enabled=True, include=["logistic", "knn"]),
        )
        assert cfg.leaderboard.enabled is True

    def test_enabled_with_non_classification_task_raises(self):
        with pytest.raises(ValidationError, match="requires model.task='classification'"):
            _minimal_config(
                model=ModelConfig(task="regression"),
                leaderboard=LeaderboardConfig(enabled=True),
            )

    def test_disabled_with_non_classification_task_does_not_raise(self):
        # Only the enabled=True combination is restricted — an unused,
        # disabled leaderboard section must never block a regression run.
        cfg = _minimal_config(
            model=ModelConfig(task="regression"),
            leaderboard=LeaderboardConfig(enabled=False),
        )
        assert cfg.leaderboard.enabled is False

    def test_unknown_include_algorithm_raises(self):
        with pytest.raises(ValidationError, match="leaderboard.include has unknown"):
            _minimal_config(leaderboard=LeaderboardConfig(enabled=True, include=["alien_net"]))

    def test_unknown_exclude_algorithm_raises(self):
        with pytest.raises(ValidationError, match="leaderboard.exclude has unknown"):
            _minimal_config(leaderboard=LeaderboardConfig(enabled=True, exclude=["alien_net"]))


# ---------------------------------------------------------------------------
# EDAConfig threshold wiring into EDAReport
# ---------------------------------------------------------------------------


class TestEDAThresholdWiring:
    def test_eda_thresholds_flow_into_eda_report(self, data_split):
        cfg = _minimal_config(
            data=DataConfig(path="dummy.parquet", target="target"),
            eda=EDAConfig(
                skewness_alert_threshold=2.5,
                zero_pct_alert_threshold=0.2,
                imbalance_alert_threshold=0.7,
                extreme_values_n=3,
                near_zero_variance_threshold=0.05,
            ),
        )
        runner = PipelineRunner(cfg)
        report = runner._run_eda(data_split)
        assert report._skewness_alert_threshold == 2.5
        assert report._zero_pct_alert_threshold == 0.2
        assert report._imbalance_alert_threshold == 0.7
        assert report._extreme_values_n == 3
        assert report._near_zero_variance_threshold == 0.05

    def test_eda_thresholds_default_to_settings_when_not_overridden(self, data_split):
        from dscompanion.config import settings

        cfg = _minimal_config(data=DataConfig(path="dummy.parquet", target="target"))
        runner = PipelineRunner(cfg)
        report = runner._run_eda(data_split)
        assert report._skewness_alert_threshold == settings.eda_skewness_alert_threshold
        assert report._extreme_values_n == settings.eda_extreme_values_n
        assert report._near_zero_variance_threshold == settings.near_zero_variance_threshold

    def test_deviation_record_has_options_default_and_user_choice(self):
        cfg = _minimal_config(eda=EDAConfig(enabled=False))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert len(deviations) == 1
        d = deviations[0]
        assert d["parameter"] == "eda.enabled"
        assert d["options"] == "True / False"
        assert d["default"] is True
        assert d["user_choice"] is False

    def test_multiple_deviations_all_detected(self):
        cfg = _minimal_config(
            eda=EDAConfig(enabled=False, bivariate=False),
            selection=SelectionConfig(remove_constant=False),
        )
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert len(deviations) == 3
        flagged = {d["parameter"] for d in deviations}
        assert flagged == {"eda.enabled", "eda.bivariate", "selection.remove_constant"}

    def test_non_monitored_field_change_not_flagged(self):
        # iv_threshold is not in the deviation checklist — changing it must not appear.
        cfg = _minimal_config(selection=SelectionConfig(iv_threshold=0.5))
        runner = PipelineRunner(cfg)
        assert runner._detect_config_deviations() == []

    def test_smote_imbalance_strategy_is_flagged(self):
        cfg = _minimal_config(target=TargetConfig(imbalance=ImbalanceConfig(strategy="smote")))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        d = next(d for d in deviations if d["parameter"] == "target.imbalance.strategy")
        assert d["default"] == "class_weight"
        assert d["user_choice"] == "smote"

    def test_tuning_enabled_is_flagged(self):
        cfg = _minimal_config(tuning=TuningConfig(enabled=True))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "tuning.enabled" for d in deviations)

    def test_vif_enabled_is_flagged(self):
        cfg = _minimal_config(selection=SelectionConfig(vif_enabled=True))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "selection.vif_enabled" for d in deviations)

    def test_multivariate_eda_is_flagged(self):
        cfg = _minimal_config(eda=EDAConfig(multivariate=True))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "eda.multivariate" for d in deviations)

    def test_lime_enabled_is_flagged(self):
        cfg = _minimal_config(explain=ExplainConfig(lime_enabled=True))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert any(d["parameter"] == "explain.lime_enabled" for d in deviations)

    def test_non_default_algorithm_is_flagged(self):
        cfg = _minimal_config(model=ModelConfig(task="classification", algorithm="logistic"))
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        d = next(d for d in deviations if d["parameter"] == "model.algorithm")
        assert d["default"] == "xgboost"
        assert d["user_choice"] == "logistic"

    def test_model_params_override_is_flagged(self):
        cfg = _minimal_config(
            model=ModelConfig(task="classification", params={"early_stopping_rounds": None})
        )
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        d = next(d for d in deviations if d["parameter"] == "model.params")
        assert d["user_choice"] == {"early_stopping_rounds": None}

    def test_empty_model_params_not_flagged(self):
        cfg = _minimal_config()
        runner = PipelineRunner(cfg)
        deviations = runner._detect_config_deviations()
        assert not any(d["parameter"] == "model.params" for d in deviations)


class TestImputerConfigWiring:
    """Regression coverage for a real bug: _build_feature_pipeline() used to

    construct FeatureProcessingPipeline() without ever passing imputer=...,
    so it always fell back to a bare SmartImputer() with every default —
    cfg.features.imputer's fields validated cleanly but had zero effect on
    the actual run.
    """

    def test_imputer_config_flows_into_feature_pipeline(self):
        cfg = _minimal_config(
            features=FeaturesConfig(
                imputer=ImputerConfig(
                    numeric_strategy="median",
                    categorical_strategy="constant",
                    fill_value="UNKNOWN",
                    add_missing_indicator=True,
                )
            )
        )
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.imputer.numeric_strategy == "median"
        assert feat_pipeline.imputer.categorical_strategy == "constant"
        assert feat_pipeline.imputer.fill_value == "UNKNOWN"
        assert feat_pipeline.imputer.add_missing_indicator is True

    def test_imputer_column_overrides_flow_into_feature_pipeline(self):
        cfg = _minimal_config(
            features=FeaturesConfig(
                imputer=ImputerConfig(
                    column_strategies={"f1": "median"}, column_fill_values={"f1": -1}
                )
            )
        )
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.imputer.column_strategies == {"f1": "median"}
        assert feat_pipeline.imputer.column_fill_values == {"f1": -1}

    def test_imputer_defaults_match_imputer_config_not_smartimputer_bare_defaults(self):
        # ImputerConfig.add_missing_indicator defaults to False, but a bare
        # SmartImputer() defaults to True — before the fix, the bare default
        # always won regardless of what cfg.features.imputer said.
        cfg = _minimal_config()
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.imputer.add_missing_indicator is False


# ---------------------------------------------------------------------------
# feature_recipes wiring (SRM.1-3) — PipelineConfig/PipelineRunner gain
# additive ColumnRecipe/FeatureTransformChain support, non-breaking to the
# ImputerConfig-based path above.
# ---------------------------------------------------------------------------


class TestFeatureRecipesConfigWiring:
    def test_feature_recipes_defaults_to_none(self):
        assert _minimal_config().feature_recipes is None

    def test_feature_recipes_accepts_column_recipe_dict(self):
        recipes = {"f1": ColumnRecipe(column="f1", steps=[TransformStep(transformer="impute")])}
        cfg = _minimal_config(feature_recipes=recipes)
        assert cfg.feature_recipes == recipes

    def test_feature_recipes_none_builds_old_smartimputer_pipeline_unchanged(self):
        # Regression guard: feature_recipes=None must leave today's path untouched.
        cfg = _minimal_config()
        feat_pipeline = PipelineRunner(cfg)._build_feature_pipeline()
        assert not isinstance(feat_pipeline, FeatureTransformChain)
        assert hasattr(feat_pipeline, "imputer")

    def test_feature_recipes_set_builds_feature_transform_chain(self):
        recipes = {"f1": ColumnRecipe(column="f1", steps=[TransformStep(transformer="impute")])}
        cfg = _minimal_config(feature_recipes=recipes)
        feat_pipeline = PipelineRunner(cfg)._build_feature_pipeline()
        assert isinstance(feat_pipeline, FeatureTransformChain)
        assert feat_pipeline.recipes == recipes

    def test_config_checks_spec_shows_feature_recipes_row_when_set(self):
        recipes = {"f1": ColumnRecipe(column="f1", steps=[TransformStep(transformer="impute")])}
        cfg = _minimal_config(feature_recipes=recipes)
        spec = PipelineRunner(cfg)._config_checks_spec()
        params = [row[0] for row in spec]
        assert "feature_recipes" in params
        assert "features.imputer.numeric_strategy" not in params

    def test_config_checks_spec_shows_imputer_rows_when_recipes_unset(self):
        cfg = _minimal_config()
        spec = PipelineRunner(cfg)._config_checks_spec()
        params = [row[0] for row in spec]
        assert "features.imputer.numeric_strategy" in params
        assert "feature_recipes" not in params

    def test_config_checks_spec_covers_previously_missing_governance_fields(self):
        """Two genuine gaps against this method's own documented contract
        ("every behavioural choice surface... not just a curated subset of
        booleans") — features.winsorizer.enabled and
        selection.remove_high_null are on/off switches in the exact same
        category as sibling flags (eda.enabled, selection.remove_constant)
        that were already tracked. Also added tuning.n_trials/metric/
        direction and leaderboard.eval_split/sort_metric — governance-
        relevant "what did this run actually optimize/rank on" facts, not
        the routine numeric thresholds (iv_threshold, correlation_threshold)
        this method deliberately excludes.
        """
        cfg = _minimal_config()
        spec = PipelineRunner(cfg)._config_checks_spec()
        params = {row[0] for row in spec}
        for expected in [
            "features.winsorizer.enabled",
            "selection.remove_high_null",
            "tuning.n_trials",
            "tuning.metric",
            "tuning.direction",
            "leaderboard.eval_split",
            "leaderboard.sort_metric",
        ]:
            assert expected in params

    def test_winsorizer_enabled_is_flagged_as_deviation(self):
        from dscompanion.pipeline.config import WinsorizerConfig

        cfg = _minimal_config(features=FeaturesConfig(winsorizer=WinsorizerConfig(enabled=True)))
        deviations = PipelineRunner(cfg)._detect_config_deviations()
        d = next(d for d in deviations if d["parameter"] == "features.winsorizer.enabled")
        assert d["default"] is False
        assert d["user_choice"] is True

    def test_remove_high_null_disabled_is_flagged_as_deviation(self):
        cfg = _minimal_config(selection=SelectionConfig(remove_high_null=False))
        deviations = PipelineRunner(cfg)._detect_config_deviations()
        d = next(d for d in deviations if d["parameter"] == "selection.remove_high_null")
        assert d["default"] is True
        assert d["user_choice"] is False

    def test_feature_recipes_row_flagged_as_deviation(self):
        recipes = {"f1": ColumnRecipe(column="f1", steps=[TransformStep(transformer="impute")])}
        cfg = _minimal_config(feature_recipes=recipes)
        deviations = PipelineRunner(cfg)._detect_config_deviations()
        assert any(d["parameter"] == "feature_recipes" for d in deviations)


class TestFeatureRecipesEndToEnd:
    """Full PipelineRunner.run() smoke test with feature_recipes set instead of
    features.imputer/encoder/scaler/distribution — proves the additive recipe
    path produces a valid trained model, mirroring TestEndToEndPipeline below.
    """

    _FEATURE_COLUMNS: list[str] = ["f1", "f4", "cat_low"]

    @pytest.fixture(scope="class")
    def synthetic_parquet(self, tmp_path_factory, synthetic_df):
        path = tmp_path_factory.mktemp("e2e_recipe_data") / "synthetic.parquet"
        synthetic_df.to_parquet(path, index=False)
        return path

    @pytest.fixture(scope="class")
    def recipe_e2e_result(self, synthetic_parquet):
        recipes = {
            "f1": ColumnRecipe(column="f1", steps=[TransformStep(transformer="impute")]),
            "f4": ColumnRecipe(column="f4", steps=[TransformStep(transformer="impute")]),
            "cat_low": ColumnRecipe(
                column="cat_low",
                steps=[
                    TransformStep(transformer="impute"),
                    TransformStep(transformer="onehot_encode"),
                ],
            ),
        }
        cfg = PipelineConfig(
            name="e2e_recipe_test",
            data=DataConfig(
                path=str(synthetic_parquet),
                target="target",
                feature_columns=self._FEATURE_COLUMNS,
            ),
            feature_recipes=recipes,
            model=ModelConfig(task="classification", algorithm="xgboost"),
            explain=ExplainConfig(shap_enabled=False),
        )
        return PipelineRunner(cfg).run()

    def test_run_returns_result_object(self, recipe_e2e_result):
        assert recipe_e2e_result is not None

    def test_model_is_fitted(self, recipe_e2e_result):
        from sklearn.utils.validation import check_is_fitted

        check_is_fitted(recipe_e2e_result.model.estimator)

    def test_metrics_dataframe_is_non_empty(self, recipe_e2e_result):
        assert isinstance(recipe_e2e_result.metrics, pd.DataFrame)
        assert len(recipe_e2e_result.metrics) > 0

    def test_feature_pipeline_is_feature_transform_chain(self, recipe_e2e_result):
        assert isinstance(recipe_e2e_result.feature_pipeline, FeatureTransformChain)

    def test_onehot_recipe_step_expanded_categorical_column(self, recipe_e2e_result):
        # cat_low is random noise relative to target, so feature *selection*
        # legitimately drops it downstream — assert against the feature
        # pipeline's own fitted output, not post-selection train_X, to test
        # FeatureTransformChain/recipe wiring independent of selection.
        output_cols = recipe_e2e_result.feature_pipeline.output_columns_["cat_low"]
        assert len(output_cols) >= 2

    def test_model_card_object_built_and_excel_capable(self, recipe_e2e_result, tmp_path):
        assert recipe_e2e_result.model_card is not None
        xlsx_path = tmp_path / "model_card.xlsx"
        recipe_e2e_result.model_card.to_excel(xlsx_path)
        assert xlsx_path.exists()
        assert xlsx_path.stat().st_size > 0


# ---------------------------------------------------------------------------
# Encoder/distribution config wiring (FT1.1/FT1.2)
# ---------------------------------------------------------------------------


class TestEncoderDistributionConfigWiring:
    """Regression coverage for a real gap: _build_feature_pipeline() used to map
    encoder.strategy="onehot" to FeatureProcessingPipeline's encoder="ordinal"
    (config validated cleanly but silently produced ordinal-encoded output,
    logged only as a warning) because no OneHotEncoder existed yet. Now that
    one does, the mapping must be real, not a fallback.
    """

    def test_onehot_strategy_no_longer_falls_back_to_ordinal(self):
        cfg = _minimal_config(features=FeaturesConfig(encoder=EncoderConfig(strategy="onehot")))
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.encoder == "onehot"

    def test_target_strategy_still_maps_to_woe(self):
        cfg = _minimal_config(features=FeaturesConfig(encoder=EncoderConfig(strategy="target")))
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.encoder == "woe"

    def test_ordinal_strategy_maps_to_ordinal(self):
        cfg = _minimal_config(features=FeaturesConfig(encoder=EncoderConfig(strategy="ordinal")))
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.encoder == "ordinal"

    def test_distribution_config_flows_into_feature_pipeline(self):
        cfg = _minimal_config(
            features=FeaturesConfig(distribution=DistributionConfig(strategy="yeo_johnson"))
        )
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.distribution == "yeo_johnson"

    def test_distribution_defaults_to_none(self):
        cfg = _minimal_config()
        runner = PipelineRunner(cfg)
        feat_pipeline = runner._build_feature_pipeline()
        assert feat_pipeline.distribution == "none"


# ---------------------------------------------------------------------------
# Full config table (report's Configuration Summary section)
# ---------------------------------------------------------------------------


class TestFullConfigTable:
    def test_default_config_all_rows_not_deviation(self):
        cfg = _minimal_config()
        runner = PipelineRunner(cfg)
        table = runner._build_full_config_table()
        assert len(table) > 0
        assert all(row["is_deviation"] is False for row in table)

    def test_table_includes_non_deviating_rows_too(self):
        # Unlike _detect_config_deviations, the full table is never filtered down —
        # a row with no override must still appear, just flagged False.
        cfg = _minimal_config(eda=EDAConfig(enabled=False))
        runner = PipelineRunner(cfg)
        table = runner._build_full_config_table()
        flagged = {row["parameter"]: row["is_deviation"] for row in table}
        assert flagged["eda.enabled"] is True
        assert flagged["eda.univariate"] is False
        assert len(table) == len(runner._config_checks_spec())

    def test_table_row_has_all_five_keys(self):
        cfg = _minimal_config()
        runner = PipelineRunner(cfg)
        row = runner._build_full_config_table()[0]
        assert set(row.keys()) == {"parameter", "choices", "default", "user_choice", "is_deviation"}

    def test_smote_and_tuning_and_vif_all_flagged_together(self):
        # Regression guard for the exact gap a user found: SMOTE + tuning.enabled +
        # vif_enabled + multivariate EDA + lime_enabled must ALL show up, not just
        # the original 7-field quality-gate subset.
        cfg = _minimal_config(
            eda=EDAConfig(multivariate=True),
            target=TargetConfig(imbalance=ImbalanceConfig(strategy="smote")),
            selection=SelectionConfig(vif_enabled=True),
            tuning=TuningConfig(enabled=True),
            explain=ExplainConfig(lime_enabled=True),
        )
        runner = PipelineRunner(cfg)
        table = runner._build_full_config_table()
        flagged = {row["parameter"] for row in table if row["is_deviation"]}
        assert flagged == {
            "eda.multivariate",
            "target.imbalance.strategy",
            "selection.vif_enabled",
            "tuning.enabled",
            "explain.lime_enabled",
        }


# ---------------------------------------------------------------------------
# _train_model — eval_set for early-stopping algorithms
# ---------------------------------------------------------------------------


class TestTrainModelEvalSet:
    """ModelFactory's xgboost classification defaults set early_stopping_rounds=20,
    which raises "Must have at least 1 validation dataset for early stopping" if
    fit() is ever called without an eval_set — independent of which split
    downstream evaluation uses. _train_model used to call model.fit(X, y) with no
    eval_set at all, so any default-config classification run would have crashed;
    caught only because experiments/bank_marketing_baseline.yaml had manually set
    early_stopping_rounds: null as a workaround, not because the bug was fixed.
    """

    @pytest.fixture
    def numeric_split(self, data_split):
        from dscompanion.split import DataSplit

        def _numeric(df):
            return df.select_dtypes(include="number").fillna(0)

        return DataSplit(
            train_X=_numeric(data_split.train_X),
            train_y=data_split.train_y,
            val_X=_numeric(data_split.val_X),
            val_y=data_split.val_y,
            test_X=_numeric(data_split.test_X),
            test_y=data_split.test_y,
            oot_X=_numeric(data_split.oot_X),
            oot_y=data_split.oot_y,
            metadata=data_split.metadata,
        )

    def test_default_xgboost_classification_does_not_crash(self, numeric_split):
        from dscompanion.models import ModelFactory

        cfg = _minimal_config(model=ModelConfig(task="classification", algorithm="xgboost"))
        runner = PipelineRunner(cfg)
        model = ModelFactory.build(task="classification", algorithm="xgboost")

        trained = runner._train_model(model, numeric_split)  # should not raise

        preds = trained.predict(numeric_split.X_test)
        assert len(preds) == len(numeric_split.X_test)

    def test_train_model_falls_back_to_test_when_val_empty(self, numeric_split):
        from dscompanion.models import ModelFactory
        from dscompanion.split import DataSplit

        split_no_val = DataSplit(
            train_X=numeric_split.train_X,
            train_y=numeric_split.train_y,
            val_X=numeric_split.train_X.iloc[:0],
            val_y=numeric_split.train_y.iloc[:0],
            test_X=numeric_split.test_X,
            test_y=numeric_split.test_y,
            oot_X=numeric_split.oot_X,
            oot_y=numeric_split.oot_y,
            metadata=numeric_split.metadata,
        )
        cfg = _minimal_config(model=ModelConfig(task="classification", algorithm="xgboost"))
        runner = PipelineRunner(cfg)
        model = ModelFactory.build(task="classification", algorithm="xgboost")

        trained = runner._train_model(model, split_no_val)  # should not raise
        preds = trained.predict(split_no_val.X_test)
        assert len(preds) == len(split_no_val.X_test)


# ---------------------------------------------------------------------------
# YAML round-trip
# ---------------------------------------------------------------------------


class TestYamlRoundTrip:
    def test_from_yaml_loads_minimal_config(self, tmp_path):
        yaml_path = tmp_path / "experiment.yaml"
        yaml_path.write_text(
            "name: bank_marketing_baseline\n"
            "data:\n"
            "  path: data/bank_marketing.parquet\n"
            "  target: subscribed\n"
            "model:\n"
            "  task: classification\n"
            "  algorithm: xgboost\n"
        )
        cfg = PipelineConfig.from_yaml(yaml_path)
        assert cfg.name == "bank_marketing_baseline"
        assert cfg.data.target == "subscribed"
        assert cfg.model.algorithm == "xgboost"
        # Defaults still populated for everything not specified
        assert cfg.split.method == "stratified"
        assert cfg.explain.shap_enabled is False

    def test_from_yaml_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            PipelineConfig.from_yaml(tmp_path / "does_not_exist.yaml")

    def test_from_yaml_invalid_field_raises_validation_error(self, tmp_path):
        yaml_path = tmp_path / "bad.yaml"
        yaml_path.write_text(
            "name: exp\n"
            "data:\n"
            "  path: dummy.parquet\n"
            "  target: y\n"
            "model:\n"
            "  task: classification\n"
            "unexpected_top_level_key: true\n"
        )
        with pytest.raises(ValidationError):
            PipelineConfig.from_yaml(yaml_path)

    def test_to_yaml_then_from_yaml_round_trip(self, tmp_path):
        original = _minimal_config(
            version="2.0",
            description="round-trip test",
            split=SplitConfig(method="random", test_size=0.3),
            model=ModelConfig(
                task="classification", algorithm="random_forest", params={"max_depth": 5}
            ),
        )
        yaml_path = tmp_path / "roundtrip.yaml"
        original.to_yaml(yaml_path)

        restored = PipelineConfig.from_yaml(yaml_path)
        assert restored == original

    def test_to_yaml_writes_all_fields_including_defaults(self, tmp_path):
        cfg = _minimal_config()
        yaml_path = tmp_path / "full.yaml"
        cfg.to_yaml(yaml_path)

        with open(yaml_path) as fh:
            raw = yaml.safe_load(fh)

        # Defaulted nested sections must be present even though never set explicitly
        assert "split" in raw
        assert "explain" in raw
        assert raw["explain"]["shap_top_n"] == 20


# ---------------------------------------------------------------------------
# _run_leaderboard — PipelineRunner's leaderboard-mode training branch
# ---------------------------------------------------------------------------


class TestRunLeaderboard:
    """PipelineConfig.leaderboard.enabled=True routes Stage 7 (train) through
    Leaderboard instead of ModelFactory.build()+_train_model() directly, and
    the run() loop updates config.model.algorithm in place to the winner so
    downstream stages (tuning, calibration, model card) see the algorithm
    that was actually trained.
    """

    @pytest.fixture
    def numeric_split(self, data_split):
        from dscompanion.split import DataSplit

        def _numeric(df):
            return df.select_dtypes(include="number").fillna(0)

        return DataSplit(
            train_X=_numeric(data_split.train_X),
            train_y=data_split.train_y,
            val_X=_numeric(data_split.val_X),
            val_y=data_split.val_y,
            test_X=_numeric(data_split.test_X),
            test_y=data_split.test_y,
            oot_X=_numeric(data_split.oot_X),
            oot_y=data_split.oot_y,
            metadata=data_split.metadata,
        )

    def test_run_leaderboard_returns_best_model_and_leaderboard_instance(self, numeric_split):
        from dscompanion.leaderboard import Leaderboard
        from dscompanion.models import ClassificationModel

        cfg = _minimal_config(
            leaderboard=LeaderboardConfig(enabled=True, include=["logistic", "decision_tree"])
        )
        runner = PipelineRunner(cfg)

        model, leaderboard = runner._run_leaderboard(numeric_split)

        assert isinstance(leaderboard, Leaderboard)
        assert isinstance(model, ClassificationModel)
        assert model is leaderboard.best_model()

    def test_run_leaderboard_respects_include_from_config(self, numeric_split):
        cfg = _minimal_config(leaderboard=LeaderboardConfig(enabled=True, include=["logistic"]))
        runner = PipelineRunner(cfg)

        _, leaderboard = runner._run_leaderboard(numeric_split)

        assert set(leaderboard.leaderboard_["algorithm"]) == {"logistic"}

    def test_run_leaderboard_winner_algorithm_matches_best_model_estimator(self, numeric_split):
        from dscompanion.models.factory import ModelFactory

        cfg = _minimal_config(
            leaderboard=LeaderboardConfig(enabled=True, include=["logistic", "decision_tree"])
        )
        runner = PipelineRunner(cfg)

        model, leaderboard = runner._run_leaderboard(numeric_split)
        winner = leaderboard.best_algorithm()

        expected_estimator = ModelFactory.build(task="classification", algorithm=winner).estimator
        assert type(model.estimator) is type(expected_estimator)


# ---------------------------------------------------------------------------
# ignore_columns (DataConfig denylist, alternative to feature_columns)
# ---------------------------------------------------------------------------


class TestIgnoreColumns:
    """DataConfig.ignore_columns — drop specific columns, keep everything else."""

    @pytest.fixture
    def uncorrelated_parquet(self, tmp_path):
        rng = np.random.RandomState(42)
        n = 200
        df = pd.DataFrame(
            {
                "x1": rng.randn(n),
                "x2": rng.uniform(0, 1, n),
                "customer_id": range(n),
                "y": [0, 1] * (n // 2),
            }
        )
        path = tmp_path / "uncorrelated.parquet"
        df.to_parquet(path, index=False)
        return path

    def _run(self, path, ignore_columns):
        cfg = PipelineConfig(
            name="ignore_columns_test",
            data={"path": str(path), "target": "y", "ignore_columns": ignore_columns},
            model={"task": "classification", "algorithm": "logistic"},
        )
        return PipelineRunner(cfg).run()

    def test_ignored_column_excluded_others_kept(self, uncorrelated_parquet):
        result = self._run(uncorrelated_parquet, ["customer_id"])
        cols = set(result.split.X_train.columns)
        assert "customer_id" not in cols
        assert {"x1", "x2"} <= cols

    def test_missing_column_raises_clear_error(self, uncorrelated_parquet):
        with pytest.raises(RuntimeError, match="ignore_columns references columns not in dataset"):
            self._run(uncorrelated_parquet, ["does_not_exist"])


# ---------------------------------------------------------------------------
# Excel data loading (DataConfig format="excel")
# ---------------------------------------------------------------------------


class TestExcelDataLoading:
    """DataConfig(format="excel") — single sheet, multi-sheet concat, and defaults."""

    @pytest.fixture
    def multi_sheet_workbook(self, tmp_path):
        df_jan = pd.DataFrame({"x1": range(50), "x2": range(50, 100), "y": [0, 1] * 25})
        df_feb = pd.DataFrame({"x1": range(100, 150), "x2": range(150, 200), "y": [0, 1] * 25})
        path = tmp_path / "multi_sheet.xlsx"
        with pd.ExcelWriter(path) as writer:
            df_jan.to_excel(writer, sheet_name="Jan", index=False)
            df_feb.to_excel(writer, sheet_name="Feb", index=False)
        return path

    def _run(self, path, sheet_name=None):
        data = {"path": str(path), "format": "excel", "target": "y"}
        if sheet_name is not None:
            data["sheet_name"] = sheet_name
        cfg = PipelineConfig(
            name="excel_test",
            data=data,
            model={"task": "classification", "algorithm": "logistic"},
        )
        return PipelineRunner(cfg).run()

    def _total_rows(self, result):
        split = result.split
        val_rows = split.X_val.shape[0] if split.X_val is not None else 0
        return split.X_train.shape[0] + split.X_test.shape[0] + val_rows

    def test_single_sheet_by_name(self, multi_sheet_workbook):
        result = self._run(multi_sheet_workbook, sheet_name="Jan")
        assert self._total_rows(result) == 50

    def test_no_sheet_name_reads_first_sheet_only(self, multi_sheet_workbook):
        result = self._run(multi_sheet_workbook)
        assert self._total_rows(result) == 50

    def test_sheet_name_list_concatenates_vertically(self, multi_sheet_workbook):
        result = self._run(multi_sheet_workbook, sheet_name=["Jan", "Feb"])
        assert self._total_rows(result) == 100

    def test_mismatched_columns_across_sheets_raises_clear_error(self, tmp_path):
        df_jan = pd.DataFrame({"x1": range(50), "x2": range(50, 100), "y": [0, 1] * 25})
        df_feb = pd.DataFrame({"x1": range(50), "x3": range(50, 100), "y": [0, 1] * 25})
        path = tmp_path / "mismatched.xlsx"
        with pd.ExcelWriter(path) as writer:
            df_jan.to_excel(writer, sheet_name="Jan", index=False)
            df_feb.to_excel(writer, sheet_name="Feb", index=False)

        with pytest.raises(RuntimeError, match="identical columns across sheets"):
            self._run(path, sheet_name=["Jan", "Feb"])


# ---------------------------------------------------------------------------
# End-to-end integration test — PipelineRunner.run() on synthetic data (#5)
# ---------------------------------------------------------------------------


class TestEndToEndPipeline:
    """Full PipelineRunner.run() smoke test — no mocking, all 13 stages execute.

    SHAP and tuning are disabled to keep runtime reasonable.  Everything else
    (EDA, imputation, encoding, scaling, selection, training, calibration,
    model card) runs at default settings on the conftest synthetic dataset.

    The pipeline is run once via a class-scoped fixture; individual tests
    each assert one aspect of the returned PipelineRunResult.
    """

    _FEATURE_COLUMNS: list[str] = [
        "f1",
        "f2",
        "f3",
        "f4",
        "f5",
        "f6",
        "f7",
        "f8",
        "f9_near_zero",
        "f10_corr_f1",
        "cat_low",
        "cat_high",
        "cat_ordinal",
        "constant_col",
    ]

    @pytest.fixture(scope="class")
    def synthetic_parquet(self, tmp_path_factory, synthetic_df):
        """Write the shared synthetic dataset to a parquet file once."""
        path = tmp_path_factory.mktemp("e2e_data") / "synthetic.parquet"
        synthetic_df.to_parquet(path, index=False)
        return path

    @pytest.fixture(scope="class")
    def e2e_result(self, synthetic_parquet):
        """Run the full pipeline once; all assertions in this class share the result."""
        cfg = PipelineConfig(
            name="e2e_test",
            data=DataConfig(
                path=str(synthetic_parquet),
                target="target",
                feature_columns=self._FEATURE_COLUMNS,
            ),
            model=ModelConfig(task="classification", algorithm="xgboost"),
            explain=ExplainConfig(shap_enabled=False),
        )
        return PipelineRunner(cfg).run()

    # ── smoke ─────────────────────────────────────────────────────────────────

    def test_run_returns_result_object(self, e2e_result):
        assert e2e_result is not None

    def test_run_produces_a_real_run_id(self, e2e_result):
        assert e2e_result.run_id is not None
        assert e2e_result.run_id != "noop-run-id"
        # yyyymmdd_hhmmss — sortable in a Workspace file browser, not uuid4().hex
        assert re.fullmatch(r"\d{8}_\d{6}(_[0-9a-f]{4})?", e2e_result.run_id)

    def test_elapsed_time_is_positive(self, e2e_result):
        assert e2e_result.elapsed_seconds > 0

    # ── metrics ───────────────────────────────────────────────────────────────

    def test_metrics_dataframe_is_non_empty(self, e2e_result):
        assert isinstance(e2e_result.metrics, pd.DataFrame)
        assert len(e2e_result.metrics) > 0

    def test_metrics_has_expected_columns(self, e2e_result):
        assert {"split", "metric", "value"} <= set(e2e_result.metrics.columns)

    def test_metrics_gini_on_test_split_is_in_unit_interval(self, e2e_result):
        row = e2e_result.metrics.query("split == 'test' and metric == 'gini'")
        assert len(row) == 1
        assert 0.0 <= float(row["value"].iloc[0]) <= 1.0

    def test_metrics_roc_auc_on_test_split_is_above_chance(self, e2e_result):
        row = e2e_result.metrics.query("split == 'test' and metric == 'roc_auc'")
        assert len(row) == 1
        # XGBoost on synthetic data with signal (f1 correlated with target) should beat 0.5
        assert float(row["value"].iloc[0]) > 0.5

    # ── report ────────────────────────────────────────────────────────────────

    def test_model_card_excel_report_is_written(self, e2e_result, tmp_path):
        xlsx_path = e2e_result.model_card.to_excel(tmp_path / "model_card.xlsx")
        assert xlsx_path.exists()
        assert xlsx_path.stat().st_size > 0

    # ── model + calibration ───────────────────────────────────────────────────

    def test_model_is_fitted(self, e2e_result):
        from sklearn.utils.validation import check_is_fitted

        check_is_fitted(e2e_result.model.estimator)

    def test_calibrator_is_fitted_for_classification(self, e2e_result):
        assert e2e_result.calibrator is not None

    def test_model_card_is_generated(self, e2e_result):
        assert e2e_result.model_card is not None

    def test_model_card_eda_split_keeps_categorical_dtypes_on_every_split(self, e2e_result):
        """Root cause of a real bug: ModelCard.split is the
        model-ready split (post impute/encode/scale/select), fully numeric
        by this point (cat_low/cat_high/cat_ordinal are ordinal-encoded),
        so Test/Validation/OOT categorical_summary() was silently empty.
        Fix: PipelineRunner now threads the pre-feature-engineering
        raw_split through as ModelCard.eda_split, which must still have
        genuine categorical dtype columns on every split — distinct from
        ModelCard.split.
        """
        card = e2e_result.model_card
        assert card.eda_split is not card.split
        # This fixture's default SplitConfig (stratified, no date_column)
        # produces an empty OOT partition — only Train/Test are populated.
        for split_X in (card.eda_split.X_train, card.eda_split.X_test):
            categorical_cols = split_X.select_dtypes(exclude="number").columns
            assert len(categorical_cols) > 0
        # card.split, by contrast, is fully numeric post-encoding.
        assert card.split.X_test.select_dtypes(exclude="number").columns.empty

    def test_to_excel_categorical_sheet_has_test_section(self, e2e_result, tmp_path):
        """End-to-end confirmation of the same fix via the actual Excel output."""
        import openpyxl

        out = e2e_result.model_card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Categorical"]
        col_b_values = [
            c.value for row in ws.iter_rows() for c in row if c.value is not None and c.column == 2
        ]
        for label in ("Train", "Test"):
            assert label in col_b_values

    # ── pipelines ─────────────────────────────────────────────────────────────

    def test_feature_pipeline_is_returned(self, e2e_result):
        assert e2e_result.feature_pipeline is not None

    def test_selection_pipeline_removes_constant_column(self, e2e_result):
        # constant_col (all zeros) must be dropped by ConstantSelector
        assert "constant_col" not in e2e_result.split.train_X.columns

    def test_selection_reduces_feature_count(self, e2e_result):
        # At minimum constant_col should be removed — final count < input count
        assert e2e_result.split.train_X.shape[1] < len(self._FEATURE_COLUMNS)

    # ── disabled optional stages ──────────────────────────────────────────────

    def test_explainer_is_none_when_shap_disabled(self, e2e_result):
        assert e2e_result.explainer is None

    def test_permutation_importance_is_none_when_disabled(self, e2e_result):
        assert e2e_result.permutation_importance is None

    def test_tuner_is_none_when_tuning_disabled(self, e2e_result):
        assert e2e_result.tuner is None

    # ── HTML report auto-write ───────────────────────────────────────────────

    def test_html_report_written_when_reporting_html_report_enabled(
        self, synthetic_parquet, tmp_path_factory
    ):
        output_dir = tmp_path_factory.mktemp("e2e_html_report")
        cfg = PipelineConfig(
            name="e2e_html_report_test",
            data=DataConfig(
                path=str(synthetic_parquet), target="target", feature_columns=self._FEATURE_COLUMNS
            ),
            model=ModelConfig(task="classification", algorithm="xgboost"),
            explain=ExplainConfig(shap_enabled=False),
            reporting=ReportingConfig(output_dir=str(output_dir), html_report=True),
        )
        result = PipelineRunner(cfg).run()
        assert result.report_path is not None
        assert Path(result.report_path).exists()
        assert Path(result.report_path).suffix == ".html"
        assert Path(result.report_path).stat().st_size > 0

    def test_html_report_not_written_by_default(self, e2e_result):
        assert e2e_result.report_path is None


# ---------------------------------------------------------------------------
# End-to-end pipeline — regression task. Mirrors TestEndToEndPipeline's classification smoke
# test — no mocking, all 13 stages execute — but with task="regression" and a
# continuous target, to catch classification-only assumptions leaking into
# the regression path (decile table, leaderboard, calibration, imbalance
# handling, and Performance Metrics diagnostics all branch on task).
# ---------------------------------------------------------------------------


class TestEndToEndPipelineRegression:
    """Full PipelineRunner.run() smoke test for task="regression" — no
    mocking, all 13 stages execute. Leaderboard/calibration/imbalance are
    all classification-only and stay disabled/no-op; SHAP is disabled to
    keep runtime reasonable, matching TestEndToEndPipeline's convention.
    """

    _FEATURE_COLUMNS: list[str] = ["f1", "f2", "f3", "f4", "cat_low", "cat_high"]

    @pytest.fixture(scope="class")
    def synthetic_regression_df(self) -> pd.DataFrame:
        """2000-row synthetic dataset with a continuous target.

        The shared conftest.py ``synthetic_df`` fixture's target is binary
        0/1 (for classification tests) and can't be reused here — this
        builds its own dataset, same style (numeric + categorical
        features), with a continuous target with known linear structure
        (f1/f3/f4 predictive, f2/cat_high pure noise) so a fitted model's r2
        should land well above chance.
        """
        rng = np.random.RandomState(7)
        n = 2000
        f1 = rng.randn(n)
        f2 = rng.randn(n)
        f3 = rng.randn(n)
        f4 = rng.randn(n)
        cat_low = rng.choice(["a", "b", "c"], size=n)
        cat_high = rng.choice([f"cat_{i}" for i in range(30)], size=n)
        target = 3.0 * f1 - 2.0 * f3 + 0.5 * f4 + rng.randn(n) * 0.5
        return pd.DataFrame(
            {
                "f1": f1,
                "f2": f2,
                "f3": f3,
                "f4": f4,
                "cat_low": cat_low,
                "cat_high": cat_high,
                "target": target,
            }
        )

    @pytest.fixture(scope="class")
    def synthetic_regression_parquet(self, tmp_path_factory, synthetic_regression_df):
        path = tmp_path_factory.mktemp("e2e_regression_data") / "synthetic_regression.parquet"
        synthetic_regression_df.to_parquet(path, index=False)
        return path

    @pytest.fixture(scope="class")
    def e2e_result(self, synthetic_regression_parquet):
        """Run the full pipeline once; all assertions in this class share the result."""
        cfg = PipelineConfig(
            name="e2e_regression_test",
            data=DataConfig(
                path=str(synthetic_regression_parquet),
                target="target",
                feature_columns=self._FEATURE_COLUMNS,
            ),
            model=ModelConfig(task="regression", algorithm="xgboost"),
            explain=ExplainConfig(shap_enabled=False),
        )
        return PipelineRunner(cfg).run()

    # ── smoke ─────────────────────────────────────────────────────────────────

    def test_run_returns_result_object(self, e2e_result):
        assert e2e_result is not None

    def test_model_is_regression_model(self, e2e_result):
        assert type(e2e_result.model).__name__ == "RegressionModel"

    # ── metrics ───────────────────────────────────────────────────────────────

    def test_metrics_has_regression_columns(self, e2e_result):
        assert {"split", "metric", "value"} <= set(e2e_result.metrics.columns)
        metrics = set(e2e_result.metrics["metric"])
        assert {"rmse", "mae", "r2", "mape", "median_absolute_error"} <= metrics

    def test_no_classification_only_metrics_present(self, e2e_result):
        metrics = set(e2e_result.metrics["metric"])
        assert not metrics & {"roc_auc", "gini", "ks_statistic", "f1", "precision", "recall"}

    def test_r2_on_test_split_is_reasonably_high(self, e2e_result):
        row = e2e_result.metrics.query("split == 'test' and metric == 'r2'")
        assert len(row) == 1
        # Strong linear signal by construction — a fitted model should recover it.
        assert float(row["value"].iloc[0]) > 0.5

    # ── classification-only stages correctly no-op for regression ──────────────

    def test_calibrator_is_none_for_regression(self, e2e_result):
        assert e2e_result.calibrator is None

    def test_leaderboard_is_none_when_disabled(self, e2e_result):
        assert e2e_result.leaderboard is None

    # ── model card ────────────────────────────────────────────────────────────

    def test_model_card_is_generated(self, e2e_result):
        assert e2e_result.model_card is not None

    def test_model_card_decile_table_is_empty_for_regression(self, e2e_result):
        # RegressionModel has no predict_proba — decile table is a
        # classification/scoring concept and is expected to be empty, not
        # crash (mirrors tests/test_docs.py's standalone unit check).
        assert e2e_result.model_card.sections_["decile_table"].empty

    def test_model_card_to_excel_does_not_crash(self, e2e_result, tmp_path):
        out = e2e_result.model_card.to_excel(tmp_path / "model_card.xlsx")
        assert out.exists()
        assert out.stat().st_size > 0

    def test_model_card_to_html_does_not_crash(self, e2e_result, tmp_path):
        out = e2e_result.model_card.to_html(tmp_path / "model_card.html")
        assert out.exists()
        assert out.stat().st_size > 0

    def test_model_card_performance_metrics_sheet_has_regression_metrics(
        self, e2e_result, tmp_path
    ):
        import openpyxl

        out = e2e_result.model_card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["Performance Metrics"]
        cell_values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        assert "rmse" in cell_values
        assert "r2" in cell_values

    def test_model_card_eda_categorical_sheet_has_test_section(self, e2e_result, tmp_path):
        """The eda_split fix (R.4-R.7's per-split EDA sections) is task-agnostic —
        confirm it still works correctly for a regression run's encoded split."""
        import openpyxl

        out = e2e_result.model_card.to_excel(tmp_path / "model_card.xlsx")
        ws = openpyxl.load_workbook(out)["EDA - Categorical"]
        col_b_values = [
            c.value for row in ws.iter_rows() for c in row if c.value is not None and c.column == 2
        ]
        assert "Train" in col_b_values
        assert "Test" in col_b_values

    # ── permutation importance (task-aware scoring) ─────────────────────────────

    def test_permutation_importance_works_for_regression(self, synthetic_regression_parquet):
        """_run_permutation_importance() already switches scoring="r2" for
        non-classification tasks — confirm that actually produces a
        populated importance table end-to-end, not just that it doesn't
        crash."""
        cfg = PipelineConfig(
            name="e2e_regression_permutation_test",
            data=DataConfig(
                path=str(synthetic_regression_parquet),
                target="target",
                feature_columns=self._FEATURE_COLUMNS,
            ),
            model=ModelConfig(task="regression", algorithm="xgboost"),
            explain=ExplainConfig(shap_enabled=False, permutation_enabled=True),
        )
        result = PipelineRunner(cfg).run()
        assert result.permutation_importance is not None
        table = result.permutation_importance.importance_table()
        assert not table.empty
        assert set(table["feature"]) <= set(self._FEATURE_COLUMNS)
        # f1/f3/f4 are predictive by construction; f2/cat_high are pure noise.
        top_feature = table.sort_values("importance_mean", ascending=False)["feature"].iloc[0]
        assert top_feature in {"f1", "f3", "f4"}


# ---------------------------------------------------------------------------
# Permutation importance stage (PI.3) — carries no SHAP/XGBoost crash risk,
# so unlike SHAP this gets a full positive end-to-end confirmation, not just
# a "disabled by default" check.
# ---------------------------------------------------------------------------


class TestPermutationImportanceStage:
    _FEATURE_COLUMNS = TestEndToEndPipeline._FEATURE_COLUMNS

    @pytest.fixture(scope="class")
    def synthetic_parquet(self, tmp_path_factory, synthetic_df):
        path = tmp_path_factory.mktemp("e2e_permutation_data") / "synthetic.parquet"
        synthetic_df.to_parquet(path, index=False)
        return path

    @pytest.fixture(scope="class")
    def e2e_result(self, synthetic_parquet):
        cfg = PipelineConfig(
            name="e2e_permutation_test",
            data=DataConfig(
                path=str(synthetic_parquet),
                target="target",
                feature_columns=self._FEATURE_COLUMNS,
            ),
            model=ModelConfig(task="classification", algorithm="xgboost"),
            explain=ExplainConfig(
                shap_enabled=False, permutation_enabled=True, permutation_n_repeats=3
            ),
        )
        return PipelineRunner(cfg).run()

    def test_permutation_importance_is_populated(self, e2e_result):
        assert e2e_result.permutation_importance is not None

    def test_importance_table_has_one_row_per_final_feature(self, e2e_result):
        table = e2e_result.permutation_importance.importance_table()
        assert len(table) == e2e_result.split.train_X.shape[1]
        assert list(table.columns) == ["feature", "importance_mean", "importance_std", "rank"]

    def test_config_deviation_is_flagged(self, e2e_result):
        assert any(
            d["parameter"] == "explain.permutation_enabled" for d in e2e_result.config_deviations
        )

    def test_explainer_still_none_when_only_permutation_enabled(self, e2e_result):
        assert e2e_result.explainer is None

    def test_model_card_explainability_section_has_permutation_table(self, e2e_result):
        section = e2e_result.model_card.sections_["explainability"]
        assert "permutation_top_features" in section
        assert "top_features" not in section


# ---------------------------------------------------------------------------
# Run output organization (RO.2) — <output_dir>/<run_id>/{model,reports,logs,eda}/
# ---------------------------------------------------------------------------


class TestRunOutputOrganization:
    _FEATURE_COLUMNS = TestEndToEndPipeline._FEATURE_COLUMNS

    @pytest.fixture(scope="class")
    def synthetic_parquet(self, tmp_path_factory, synthetic_df):
        path = tmp_path_factory.mktemp("ro_data") / "synthetic.parquet"
        synthetic_df.to_parquet(path, index=False)
        return path

    @pytest.fixture(scope="class")
    def output_root(self, tmp_path_factory):
        return tmp_path_factory.mktemp("ro_output_root")

    @pytest.fixture(scope="class")
    def e2e_result(self, synthetic_parquet, output_root):
        cfg = PipelineConfig(
            name="ro_test",
            data=DataConfig(
                path=str(synthetic_parquet),
                target="target",
                feature_columns=self._FEATURE_COLUMNS,
            ),
            model=ModelConfig(task="classification", algorithm="xgboost"),
            explain=ExplainConfig(shap_enabled=False),
            reporting=ReportingConfig(output_dir=str(output_root)),
        )
        return PipelineRunner(cfg).run()

    def test_run_dir_is_populated(self, e2e_result):
        assert e2e_result.run_dir is not None

    def test_run_dir_name_matches_run_id(self, e2e_result):
        assert Path(e2e_result.run_dir).name == e2e_result.run_id

    def test_run_dir_is_nested_under_output_dir(self, e2e_result, output_root):
        assert Path(e2e_result.run_dir).parent == Path(output_root)

    @pytest.mark.parametrize("subdir", ["model", "reports", "logs", "eda"])
    def test_run_dir_has_expected_subdirectory(self, e2e_result, subdir):
        assert (Path(e2e_result.run_dir) / subdir).is_dir()

    def test_log_path_is_populated_and_non_empty(self, e2e_result):
        assert e2e_result.log_path is not None
        assert Path(e2e_result.log_path).is_file()
        assert Path(e2e_result.log_path).stat().st_size > 0

    def test_log_path_is_inside_run_dir_logs_subdirectory(self, e2e_result):
        assert Path(e2e_result.log_path).parent == Path(e2e_result.run_dir) / "logs"

    def test_log_file_contains_pipeline_stage_lines(self, e2e_result):
        content = Path(e2e_result.log_path).read_text()
        assert "Loading data" in content
        assert "Pipeline complete" in content

    def test_model_path_is_populated_and_loadable(self, e2e_result):
        from dscompanion.models.base import BaseDSCompanionModel

        assert e2e_result.model_path is not None
        assert Path(e2e_result.model_path).is_file()
        loaded = BaseDSCompanionModel.load(e2e_result.model_path)
        assert type(loaded) is type(e2e_result.model)

    def test_model_path_is_inside_run_dir_model_subdirectory(self, e2e_result):
        assert Path(e2e_result.model_path).parent == Path(e2e_result.run_dir) / "model"

    def test_excel_report_path_is_populated_and_non_empty(self, e2e_result):
        assert e2e_result.excel_report_path is not None
        assert Path(e2e_result.excel_report_path).is_file()
        assert Path(e2e_result.excel_report_path).stat().st_size > 0

    def test_excel_report_path_is_inside_run_dir_reports_subdirectory(self, e2e_result):
        assert Path(e2e_result.excel_report_path).parent == Path(e2e_result.run_dir) / "reports"

    def test_config_yaml_saved_directly_in_run_dir(self, e2e_result):
        config_path = Path(e2e_result.run_dir) / "config.yaml"
        assert config_path.is_file()
        loaded = yaml.safe_load(config_path.read_text())
        assert loaded["name"] == "ro_test"
        assert loaded["model"]["algorithm"] == "xgboost"

    def test_log_capture_works_even_when_ambient_root_logger_is_at_warning(
        self, synthetic_parquet, output_root
    ):
        # Regression test for a confirmed Databricks bug (2026-09-10): a notebook
        # host pre-installs its own root logger handler *before* `import dscompanion`
        # runs, which silently makes dscompanion/__init__.py's own
        # logging.basicConfig(level=INFO) a no-op (per Python's own docs —
        # basicConfig() does nothing if the root logger already has handlers).
        # Net effect: dscompanion's effective level stays at root's default WARNING,
        # so logger.info() calls never even create a LogRecord — no FileHandler,
        # however attached, can capture what was never created.
        # _attach_run_log_handler() must not depend on ambient logging
        # configuration it doesn't control.
        root_logger = logging.getLogger()
        dscompanion_logger = logging.getLogger("dscompanion")
        prev_root_level = root_logger.level
        prev_dscompanion_level = dscompanion_logger.level
        prev_root_handlers = list(root_logger.handlers)
        try:
            for h in prev_root_handlers:
                root_logger.removeHandler(h)
            root_logger.addHandler(logging.StreamHandler())
            root_logger.setLevel(logging.WARNING)
            dscompanion_logger.setLevel(logging.NOTSET)

            cfg = PipelineConfig(
                name="hostile_logging_test",
                data=DataConfig(
                    path=str(synthetic_parquet),
                    target="target",
                    feature_columns=self._FEATURE_COLUMNS,
                ),
                model=ModelConfig(task="classification", algorithm="xgboost"),
                explain=ExplainConfig(shap_enabled=False),
                reporting=ReportingConfig(output_dir=str(output_root / "hostile")),
            )
            result = PipelineRunner(cfg).run()
            content = Path(result.log_path).read_text()
            assert "Loading data" in content
            assert "Pipeline complete" in content
        finally:
            for h in list(root_logger.handlers):
                root_logger.removeHandler(h)
            for h in prev_root_handlers:
                root_logger.addHandler(h)
            root_logger.setLevel(prev_root_level)
            dscompanion_logger.setLevel(prev_dscompanion_level)

    def test_no_handler_leaked_onto_dscompanion_logger_after_run(self, e2e_result):
        # Two runs already happened by the time this fixture is shared (this class's
        # e2e_result plus every earlier class's own runs in the same pytest session) —
        # if the FileHandler were never removed, this count would keep growing.
        dscompanion_logger = logging.getLogger("dscompanion")
        file_handlers = [
            h for h in dscompanion_logger.handlers if isinstance(h, logging.FileHandler)
        ]
        assert file_handlers == []


class TestReportingConfigHtmlReport:
    def test_html_report_defaults_to_false(self):
        assert ReportingConfig().html_report is False

    def test_output_dir_defaults_to_reports(self):
        assert ReportingConfig().output_dir == "./reports"

    def test_html_report_and_output_dir_are_settable(self):
        cfg = ReportingConfig(html_report=True, output_dir="/tmp/custom")
        assert cfg.html_report is True
        assert cfg.output_dir == "/tmp/custom"
